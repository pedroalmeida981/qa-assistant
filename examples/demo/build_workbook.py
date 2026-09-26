"""
Builds the initial workbook of the demo workspace (plan, task 2.1).

Public practice site: https://www.saucedemo.com/ (demo accounts published by the
site itself; the credentials live only in the local .env).
Cases approved by the user on 23.09.2026 (methodology section 5.3).
Expected results = the site's normal published behaviour (standard_user), with
the interface messages in English.

API area (plan 3.1, technique-api.md): public practice API
https://restful-booker.herokuapp.com, expected results from the published
documentation (/apidoc); where the documentation is silent, the case cites the
HTTP standard (RFC 9110). Cases proposed on 24.09.2026 and approved by delegation
(the user asked to go ahead without asking): review.

Accessibility area (plan 3.2, technique-accessibility.md): WCAG 2.1 AA;
automated cases (axe-core, scripts/audit_a11y.py) and guided manual ones.
Also approved by delegation on 24.09.2026: review.

Mobile area (plan 3.3, technique-mobile.md): Maestro flows in flows/, run in
Maestro Web (beta) because there is no device on this machine; expected = the
site's normal behaviour. Approved by delegation on 24.09.2026: review.

    python examples/demo/build_workbook.py [target.xlsx]
"""
from __future__ import annotations

import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

STATUSES = ["Not run", "Pass", "Fail", "Blocked", "N/A"]
CASE_HEADERS = ["ID", "Title", "Role", "Technique", "Priority", "Preconditions", "Steps",
                "Expected Result", "Status", "Actual Result", "Notes"]
LOGIN = "Login page open, no session started."
SESSION = "Session started with the given role; Products page open; cart empty."
TOKEN = "Admin token obtained with POST /auth (TC-API-002), saved as a session secret."
CREATED = "Test booking created in TC-API-003 (firstname with the QA-AUTO- prefix)."
AXE = "Automated audit (python scripts/audit_a11y.py), page declared in accessibility.pages."
AXE_EXP = "No WCAG 2.1 A/AA violations detected by axe-core (best practices and incomplete results do not count)."
MAESTRO = "Maestro flow in flows/ (python scripts/run_mobile.py); no session; credentials from .env as MAESTRO_<VAR>."

CASES = {
    "Login": [
        ("TC-LOGIN-001", "Login with valid credentials", "standard_user", "Equivalence partitioning", "High", LOGIN,
         "1. Fill in Username and Password with the role's account.\n2. Click Login.",
         "The \"Products\" page opens (/inventory.html)."),
        ("TC-LOGIN-002", "Login of a locked out user", "locked_out_user", "Role matrix", "High", LOGIN,
         "1. Fill in Username and Password with the role's account.\n2. Click Login.",
         "Shows the error \"Epic sadface: Sorry, this user has been locked out.\" and stays on the login page."),
        ("TC-LOGIN-003", "Login without username", "—", "Equivalence partitioning", "Medium", LOGIN,
         "1. Leave Username empty and fill in Password.\n2. Click Login.",
         "Shows the error \"Epic sadface: Username is required\"."),
        ("TC-LOGIN-004", "Login without password", "standard_user", "Equivalence partitioning", "Medium", LOGIN,
         "1. Fill in Username with the role's account and leave Password empty.\n2. Click Login.",
         "Shows the error \"Epic sadface: Password is required\"."),
        ("TC-LOGIN-005", "Login with a wrong password", "standard_user", "Equivalence partitioning", "Medium", LOGIN,
         "1. Fill in Username with the role's account and Password with a wrong value.\n2. Click Login.",
         "Shows the error \"Epic sadface: Username and password do not match any user in this service\"."),
        ("TC-LOGIN-006", "Direct access to the inventory without a session", "—", "Error guessing", "High",
         "No session started.",
         "1. Open /inventory.html directly.",
         ("Shows the login form with the error \"Epic sadface: You can only access '/inventory.html' when you "
          "are logged in.\" (the URL may stay /inventory.html — user decision, 24.09.2026).")),
        ("TC-LOGIN-007", "Logout", "standard_user", "State transition", "Medium", SESSION,
         "1. Open the side menu.\n2. Click Logout.\n3. Use the browser's Back.",
         "Goes back to the login page; Back does not reopen the inventory without a new session."),
    ],
    "Products": [
        ("TC-PROD-001", "Product list", "standard_user", "Equivalence partitioning", "High", SESSION,
         "1. Look at the product list.",
         "6 products are shown, each with name, description, price, image and an \"Add to cart\" button."),
        ("TC-PROD-002", "Sort by price (low to high)", "standard_user", "Equivalence partitioning", "Medium", SESSION,
         "1. Choose \"Price (low to high)\" in the sort selector.",
         "The products are in ascending order of price."),
        ("TC-PROD-003", "Sort by name (Z to A)", "standard_user", "Equivalence partitioning", "Low", SESSION,
         "1. Choose \"Name (Z to A)\" in the sort selector.",
         "The products are in descending alphabetical order of name."),
        ("TC-PROD-004", "Product detail", "standard_user", "Equivalence partitioning", "Medium", SESSION,
         "1. Note the name and price of the first product.\n2. Click the product's name.\n3. Click \"Back to products\".",
         "The detail shows the same name and price as the list; \"Back to products\" goes back to the list."),
        ("TC-PROD-005", "Product images (problem_user)", "problem_user", "Role matrix", "Medium", SESSION,
         "1. Look at each product's image and compare it with standard_user's.",
         "Each product shows its own image, the same one shown to standard_user."),
    ],
    "Checkout": [
        ("TC-CHK-001", "Add to cart", "standard_user", "State transition", "High", SESSION,
         "1. Click \"Add to cart\" on the first product.",
         "The cart counter shows 1 and the product's button becomes \"Remove\"."),
        ("TC-CHK-002", "Remove from the cart", "standard_user", "State transition", "Medium",
         SESSION + " One product in the cart.",
         "1. Open the cart.\n2. Click \"Remove\" on the product.",
         "The cart becomes empty and the cart counter disappears."),
        ("TC-CHK-003", "Checkout without First Name", "standard_user", "Equivalence partitioning", "Medium",
         SESSION + " One product in the cart.",
         "1. Open the cart and click Checkout.\n2. Fill in Last Name and Postal Code, leave First Name empty.\n3. Click Continue.",
         "Shows the error \"Error: First Name is required\"."),
        ("TC-CHK-004", "Checkout without Postal Code", "standard_user", "Equivalence partitioning", "Medium",
         SESSION + " One product in the cart.",
         "1. Open the cart and click Checkout.\n2. Fill in First Name and Last Name, leave Postal Code empty.\n3. Click Continue.",
         "Shows the error \"Error: Postal Code is required\"."),
        ("TC-CHK-005", "Total in the checkout overview", "standard_user", "Boundary values (calculation)", "High",
         SESSION + " Two products in the cart.",
         "1. Open the cart and click Checkout.\n2. Fill in the three fields with fictional data.\n3. Click Continue.",
         "Total = Item total + Tax, and Item total = sum of the prices of the products in the cart."),
        ("TC-CHK-006", "Complete the purchase", "standard_user", "State transition", "High",
         SESSION + " One product in the cart.",
         "1. Do the checkout with fictional data up to the overview.\n2. Click Finish.",
         "Shows \"Thank you for your order!\" and the cart becomes empty."),
    ],
    "API": [
        ("TC-API-001", "Health check", "—", "Equivalence partitioning", "High", "API reachable.",
         "1. GET /ping.",
         "Answers 201 Created (HealthCheck in the documentation)."),
        ("TC-API-002", "Create a token with valid credentials", "admin (API)", "Equivalence partitioning", "High",
         "Admin credentials in .env.",
         "1. POST /auth with the admin username and password (JSON).",
         "Answers 200 with a non-empty \"token\" field (CreateToken)."),
        ("TC-API-003", "Create a valid booking", "—", "Equivalence partitioning", "High", "API reachable.",
         "1. POST /booking with firstname QA-AUTO-…, lastname, totalprice, depositpaid, bookingdates and additionalneeds.",
         "Answers 200 with a numeric bookingid and a booking equal to the one sent (CreateBooking)."),
        ("TC-API-004", "Read the booking created", "—", "State transition", "High", CREATED,
         "1. GET /booking/{id}.",
         "Answers 200 with the same data sent at creation (GetBooking)."),
        ("TC-API-005", "Filter bookings by firstname", "—", "Equivalence partitioning", "Medium", CREATED,
         "1. GET /booking?firstname=<firstname of the booking created>.",
         "Answers 200 with a list that contains the bookingid created (GetBookingIds)."),
        ("TC-API-006", "Update without authentication", "—", "Role matrix", "High", CREATED,
         "1. PUT /booking/{id} with new data, without Cookie nor Authorization.\n2. GET /booking/{id}.",
         "The request is refused (4xx: the documentation requires a token or Basic on PUT) and the booking stays the same."),
        ("TC-API-007", "Update with a token", "admin (API)", "State transition", "High", CREATED + " " + TOKEN,
         "1. PUT /booking/{id} with new data and Cookie: token=<token>.\n2. GET /booking/{id}.",
         "Answers 200 with the new data, and the next GET returns it (UpdateBooking)."),
        ("TC-API-008", "Partial update with a token", "admin (API)", "Equivalence partitioning", "Medium",
         CREATED + " " + TOKEN,
         "1. PATCH /booking/{id} with only a new firstname and Cookie: token=<token>.\n2. GET /booking/{id}.",
         "Answers 200; only the firstname changes, the other fields stay the same (PartialUpdateBooking)."),
        ("TC-API-009", "Create a booking without required fields", "—", "Equivalence partitioning", "Medium",
         "API reachable.",
         "1. POST /booking with JSON with only lastname.",
         ("Answers 4xx (client error), never 5xx. Source: the documentation is silent about errors; "
          "RFC 9110 §15.5–15.6 (invalid input = 4xx).")),
        ("TC-API-010", "Delete the booking with a token", "admin (API)", "State transition", "Medium",
         CREATED + " " + TOKEN,
         "1. DELETE /booking/{id} with Cookie: token=<token>.",
         "Answers 201 Created (DeleteBooking in the documentation)."),
    ],
    "Accessibility": [
        ("TC-A11Y-001", "Login: automated WCAG 2.1 AA", "—", "Error guessing (axe-core)", "High", AXE,
         "1. Audit the login page.", AXE_EXP),
        ("TC-A11Y-002", "Products: automated WCAG 2.1 AA", "standard_user", "Error guessing (axe-core)", "High",
         AXE + " With login.", "1. Audit the products page.", AXE_EXP),
        ("TC-A11Y-003", "Cart: automated WCAG 2.1 AA", "standard_user", "Error guessing (axe-core)", "Medium",
         AXE + " With login.", "1. Audit the cart page.", AXE_EXP),
        ("TC-A11Y-004", "Checkout (information): automated WCAG 2.1 AA", "standard_user", "Error guessing (axe-core)",
         "Medium", AXE + " With login.", "1. Audit the checkout-information page.", AXE_EXP),
        ("TC-A11Y-005", "Contrast of the elements axe did not decide", "standard_user", "Manual checklist (WCAG)",
         "Medium", "Incomplete color-contrast results of the Products audit.",
         "1. Measure the contrast of each element listed by axe as incomplete (real text and background).",
         "Contrast of at least 4.5:1 (normal text) or 3:1 (large text) — criterion 1.4.3 AA."),
        ("TC-A11Y-006", "Each page's title identifies its content", "standard_user", "Manual checklist (WCAG)",
         "Medium", "Pages login, products, cart and checkout-information.",
         "1. Read the title of each page.",
         ("Each page has a title that identifies its content or purpose, different between different pages "
          "— criterion 2.4.2 A (Understanding 2.4.2, technique G88).")),
        ("TC-A11Y-007", "Login operable with the keyboard only", "standard_user", "Manual checklist (WCAG)", "High",
         "Login page open, no session started.",
         "1. With Tab, reach Username, Password and Login.\n2. Fill in and submit with Enter.",
         "The login is done with the keyboard only, with a visible focus on each field — criteria 2.1.1 A and 2.4.7 AA."),
    ],
    "Mobile": [
        ("TC-MOB-001", "Login with valid credentials (flow)", "standard_user", "Equivalence partitioning", "High",
         MAESTRO, "1. Run flows/TC-MOB-001_login.yaml.",
         "After the login the \"Products\" page appears."),
        ("TC-MOB-002", "Login of a locked out user (flow)", "locked_out_user", "Role matrix", "High",
         MAESTRO, "1. Run flows/TC-MOB-002_locked.yaml.",
         "Shows \"Epic sadface: Sorry, this user has been locked out.\" and does not log in."),
        ("TC-MOB-003", "Add to cart (flow)", "standard_user", "State transition", "Medium",
         MAESTRO, "1. Run flows/TC-MOB-003_cart.yaml.",
         "The product's button becomes \"Remove\" and the cart counter shows 1."),
    ],
}

DEFECT_HEADERS = ["ID", "Title", "Description", "Area", "Case", "Role", "Severity", "Priority",
                  "Status", "Reported by", "Date"]
VIEW_HEADERS = ["ID", "Priority", "Severity", "Status", "Area", "Title", "Case", "Role"]
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")


def _header(ws, row, headers):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row, i, h)
        c.font, c.fill = HEAD, HEAD_FILL


def build(path: str) -> str:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Summary")
    ws["A1"] = "QA demo — saucedemo.com (public practice site)"
    ws["A1"].font = Font(bold=True, size=14)
    _header(ws, 3, ["Area", "Total"] + STATUSES + ["% Run"])
    for i, sheet in enumerate(CASES, start=4):
        ws.cell(i, 1, sheet)
        ws.cell(i, 2, f"=COUNTA('{sheet}'!A2:A200)")
        for j, e in enumerate(STATUSES, start=3):
            ws.cell(i, j, f"=COUNTIF('{sheet}'!I2:I200,\"{e}\")")
        ws.cell(i, 8, f"=IFERROR((B{i}-C{i})/B{i},0)").number_format = "0%"
    t = 4 + len(CASES)
    ws.cell(t, 1, "Total").font = Font(bold=True)
    for col in "BCDEFG":
        ws[f"{col}{t}"] = f"=SUM({col}4:{col}{t - 1})"
    ws.cell(t, 8, f"=IFERROR((B{t}-C{t})/B{t},0)").number_format = "0%"
    ws.cell(t + 2, 1, "% Run = cases outside Not run (N/A and Blocked count as run).")

    for sheet, cases in CASES.items():
        ws = wb.create_sheet(sheet)
        _header(ws, 1, CASE_HEADERS)
        for c in cases:
            ws.append(list(c) + ["Not run", None, None])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add("I2:I200")
        for col, w in zip("ABCDEFGHIJK", (14, 32, 16, 22, 10, 30, 45, 45, 12, 35, 35)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "C2"

    ws = wb.create_sheet("Defects")
    _header(ws, 1, DEFECT_HEADERS)
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Active Defects")
    ws["A1"] = "Active defects (regenerated by the adapter — do not edit by hand)"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "no defects recorded"
    _header(ws, 3, VIEW_HEADERS)

    wb.save(path)
    return path


if __name__ == "__main__":
    for _s in (sys.stdout, sys.stderr):  # output in UTF-8 also through a pipe on Windows
        _s.reconfigure(encoding="utf-8", errors="replace")
    here = os.path.dirname(os.path.abspath(__file__))
    dest = sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "workbook", "demo-plan.xlsx")
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    print(build(dest))
