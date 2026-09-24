"""
Gera o workbook inicial do workspace de demonstracao (plano, tarefa 2.1).

Site de pratica publico: https://www.saucedemo.com/ (contas de demonstracao
publicadas pelo proprio site; as credenciais vivem so no .env local).
Casos aprovados pelo utilizador em 23.09.2026 (metodologia-core seccao 5.3).
Resultados esperados = comportamento normal publicado do site (standard_user),
com as mensagens da interface em ingles.

Area API (plano 3.1, tecnica-api.md): API publica de pratica
https://restful-booker.herokuapp.com, esperados da documentacao publicada
(/apidoc); onde a documentacao e omissa, o caso cita a norma HTTP (RFC 9110).
Casos propostos a 24.09.2026 e aprovados por delegacao (o utilizador pediu
para avancar sem perguntar): rever.

    python examples/demo/gerar_workbook.py [destino.xlsx]
"""
from __future__ import annotations

import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ESTADOS = ["Por testar", "Pass", "Fail", "Blocked", "N/A"]
CASE_HEADERS = ["ID", "Título", "Role", "Técnica", "Prioridade", "Pré-condições", "Passos",
                "Resultado Esperado", "Estado", "Resultado Obtido", "Observações"]
LOGIN = "Página de login aberta, sem sessão iniciada."
SESSAO = "Sessão iniciada com a role indicada; página Products aberta; carrinho vazio."
TOKEN = "Token de admin obtido com POST /auth (CT-API-002), guardado como segredo da sessão."
CRIADA = "Reserva de teste criada em CT-API-003 (firstname com prefixo QA-AUTO-)."

CASOS = {
    "Login": [
        ("CT-LOGIN-001", "Login com credenciais válidas", "standard_user", "Partição de equivalência", "Alta", LOGIN,
         "1. Preencher Username e Password com a conta da role.\n2. Clicar em Login.",
         "Abre a página \"Products\" (/inventory.html)."),
        ("CT-LOGIN-002", "Login de utilizador bloqueado", "locked_out_user", "Matriz de roles", "Alta", LOGIN,
         "1. Preencher Username e Password com a conta da role.\n2. Clicar em Login.",
         "Mostra o erro \"Epic sadface: Sorry, this user has been locked out.\" e fica na página de login."),
        ("CT-LOGIN-003", "Login sem username", "—", "Partição de equivalência", "Média", LOGIN,
         "1. Deixar Username vazio e preencher Password.\n2. Clicar em Login.",
         "Mostra o erro \"Epic sadface: Username is required\"."),
        ("CT-LOGIN-004", "Login sem password", "standard_user", "Partição de equivalência", "Média", LOGIN,
         "1. Preencher Username com a conta da role e deixar Password vazio.\n2. Clicar em Login.",
         "Mostra o erro \"Epic sadface: Password is required\"."),
        ("CT-LOGIN-005", "Login com password errada", "standard_user", "Partição de equivalência", "Média", LOGIN,
         "1. Preencher Username com a conta da role e Password com um valor errado.\n2. Clicar em Login.",
         "Mostra o erro \"Epic sadface: Username and password do not match any user in this service\"."),
        ("CT-LOGIN-006", "Acesso direto ao inventário sem sessão", "—", "Error guessing", "Alta",
         "Sem sessão iniciada.",
         "1. Abrir diretamente /inventory.html.",
         "Volta à página de login com o erro \"Epic sadface: You can only access '/inventory.html' when you are logged in.\"."),
        ("CT-LOGIN-007", "Logout", "standard_user", "Transição de estados", "Média", SESSAO,
         "1. Abrir o menu lateral.\n2. Clicar em Logout.\n3. Usar o Back do browser.",
         "Volta à página de login; o Back não reabre o inventário sem nova sessão."),
    ],
    "Produtos": [
        ("CT-PROD-001", "Lista de produtos", "standard_user", "Partição de equivalência", "Alta", SESSAO,
         "1. Observar a lista de produtos.",
         "São apresentados 6 produtos, cada um com nome, descrição, preço, imagem e botão \"Add to cart\"."),
        ("CT-PROD-002", "Ordenar por preço (low to high)", "standard_user", "Partição de equivalência", "Média", SESSAO,
         "1. Escolher \"Price (low to high)\" no seletor de ordenação.",
         "Os produtos ficam por ordem crescente de preço."),
        ("CT-PROD-003", "Ordenar por nome (Z to A)", "standard_user", "Partição de equivalência", "Baixa", SESSAO,
         "1. Escolher \"Name (Z to A)\" no seletor de ordenação.",
         "Os produtos ficam por ordem alfabética decrescente de nome."),
        ("CT-PROD-004", "Detalhe do produto", "standard_user", "Partição de equivalência", "Média", SESSAO,
         "1. Anotar o nome e o preço do primeiro produto.\n2. Clicar no nome do produto.\n3. Clicar em \"Back to products\".",
         "O detalhe mostra o mesmo nome e preço da lista; \"Back to products\" volta à lista."),
        ("CT-PROD-005", "Imagens dos produtos (problem_user)", "problem_user", "Matriz de roles", "Média", SESSAO,
         "1. Observar a imagem de cada produto e comparar com a do standard_user.",
         "Cada produto mostra a sua própria imagem, igual à apresentada ao standard_user."),
    ],
    "Checkout": [
        ("CT-CHK-001", "Adicionar ao carrinho", "standard_user", "Transição de estados", "Alta", SESSAO,
         "1. Clicar em \"Add to cart\" no primeiro produto.",
         "O contador do carrinho mostra 1 e o botão do produto passa a \"Remove\"."),
        ("CT-CHK-002", "Remover do carrinho", "standard_user", "Transição de estados", "Média",
         SESSAO + " Um produto no carrinho.",
         "1. Abrir o carrinho.\n2. Clicar em \"Remove\" no produto.",
         "O carrinho fica vazio e o contador do carrinho desaparece."),
        ("CT-CHK-003", "Checkout sem First Name", "standard_user", "Partição de equivalência", "Média",
         SESSAO + " Um produto no carrinho.",
         "1. Abrir o carrinho e clicar em Checkout.\n2. Preencher Last Name e Postal Code, deixar First Name vazio.\n3. Clicar em Continue.",
         "Mostra o erro \"Error: First Name is required\"."),
        ("CT-CHK-004", "Checkout sem Postal Code", "standard_user", "Partição de equivalência", "Média",
         SESSAO + " Um produto no carrinho.",
         "1. Abrir o carrinho e clicar em Checkout.\n2. Preencher First Name e Last Name, deixar Postal Code vazio.\n3. Clicar em Continue.",
         "Mostra o erro \"Error: Postal Code is required\"."),
        ("CT-CHK-005", "Total no resumo do checkout", "standard_user", "Valores-limite (cálculo)", "Alta",
         SESSAO + " Dois produtos no carrinho.",
         "1. Abrir o carrinho e clicar em Checkout.\n2. Preencher os três campos com dados fictícios.\n3. Clicar em Continue.",
         "Total = Item total + Tax, e Item total = soma dos preços dos produtos no carrinho."),
        ("CT-CHK-006", "Concluir a compra", "standard_user", "Transição de estados", "Alta",
         SESSAO + " Um produto no carrinho.",
         "1. Fazer o checkout com dados fictícios até ao resumo.\n2. Clicar em Finish.",
         "Mostra \"Thank you for your order!\" e o carrinho fica vazio."),
    ],
    "API": [
        ("CT-API-001", "Health check", "—", "Partição de equivalência", "Alta", "API acessível.",
         "1. GET /ping.",
         "Responde 201 Created (HealthCheck na documentação)."),
        ("CT-API-002", "Criar token com credenciais válidas", "admin (API)", "Partição de equivalência", "Alta",
         "Credenciais de admin no .env.",
         "1. POST /auth com username e password de admin (JSON).",
         "Responde 200 com um campo \"token\" não vazio (CreateToken)."),
        ("CT-API-003", "Criar reserva válida", "—", "Partição de equivalência", "Alta", "API acessível.",
         "1. POST /booking com firstname QA-AUTO-…, lastname, totalprice, depositpaid, bookingdates e additionalneeds.",
         "Responde 200 com bookingid numérico e booking igual ao enviado (CreateBooking)."),
        ("CT-API-004", "Ler a reserva criada", "—", "Transição de estados", "Alta", CRIADA,
         "1. GET /booking/{id}.",
         "Responde 200 com os mesmos dados enviados na criação (GetBooking)."),
        ("CT-API-005", "Filtrar reservas por firstname", "—", "Partição de equivalência", "Média", CRIADA,
         "1. GET /booking?firstname=<firstname da reserva criada>.",
         "Responde 200 com uma lista que contém o bookingid criado (GetBookingIds)."),
        ("CT-API-006", "Atualizar sem autenticação", "—", "Matriz de roles", "Alta", CRIADA,
         "1. PUT /booking/{id} com dados novos, sem Cookie nem Authorization.\n2. GET /booking/{id}.",
         "O pedido é recusado (4xx: a documentação exige token ou Basic no PUT) e a reserva fica igual."),
        ("CT-API-007", "Atualizar com token", "admin (API)", "Transição de estados", "Alta", CRIADA + " " + TOKEN,
         "1. PUT /booking/{id} com dados novos e Cookie: token=<token>.\n2. GET /booking/{id}.",
         "Responde 200 com os dados novos, e o GET seguinte devolve-os (UpdateBooking)."),
        ("CT-API-008", "Atualização parcial com token", "admin (API)", "Partição de equivalência", "Média",
         CRIADA + " " + TOKEN,
         "1. PATCH /booking/{id} só com firstname novo e Cookie: token=<token>.\n2. GET /booking/{id}.",
         "Responde 200; só o firstname muda, os outros campos ficam iguais (PartialUpdateBooking)."),
        ("CT-API-009", "Criar reserva sem campos obrigatórios", "—", "Partição de equivalência", "Média",
         "API acessível.",
         "1. POST /booking com JSON só com lastname.",
         ("Responde 4xx (erro do cliente), nunca 5xx. Fonte: a documentação é omissa sobre erros; "
          "RFC 9110 §15.5–15.6 (entrada inválida = 4xx).")),
        ("CT-API-010", "Apagar a reserva com token", "admin (API)", "Transição de estados", "Média",
         CRIADA + " " + TOKEN,
         "1. DELETE /booking/{id} com Cookie: token=<token>.",
         "Responde 201 Created (DeleteBooking na documentação)."),
    ],
}

DEFECT_HEADERS = ["ID", "Título", "Descrição", "Área", "Caso", "Role", "Severidade", "Prioridade",
                  "Estado", "Reportado por", "Data"]
VIEW_HEADERS = ["ID", "Prioridade", "Severidade", "Estado", "Área", "Título", "Caso", "Role"]
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")


def _header(ws, row, headers):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row, i, h)
        c.font, c.fill = HEAD, HEAD_FILL


def gerar(path: str) -> str:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Resumo")
    ws["A1"] = "Demo QA — saucedemo.com (site público de prática)"
    ws["A1"].font = Font(bold=True, size=14)
    _header(ws, 3, ["Área", "Total"] + ESTADOS + ["% Executado"])
    for i, sheet in enumerate(CASOS, start=4):
        ws.cell(i, 1, sheet)
        ws.cell(i, 2, f"=COUNTA('{sheet}'!A2:A200)")
        for j, e in enumerate(ESTADOS, start=3):
            ws.cell(i, j, f"=COUNTIF('{sheet}'!I2:I200,\"{e}\")")
        ws.cell(i, 8, f"=IFERROR((B{i}-C{i})/B{i},0)").number_format = "0%"
    t = 4 + len(CASOS)
    ws.cell(t, 1, "Total").font = Font(bold=True)
    for col in "BCDEFG":
        ws[f"{col}{t}"] = f"=SUM({col}4:{col}{t - 1})"
    ws.cell(t, 8, f"=IFERROR((B{t}-C{t})/B{t},0)").number_format = "0%"
    ws.cell(t + 2, 1, "% Executado = casos fora de Por testar (N/A e Blocked contam como executados).")

    for sheet, casos in CASOS.items():
        ws = wb.create_sheet(sheet)
        _header(ws, 1, CASE_HEADERS)
        for c in casos:
            ws.append(list(c) + ["Por testar", None, None])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        dv = DataValidation(type="list", formula1='"' + ",".join(ESTADOS) + '"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add("I2:I200")
        for col, w in zip("ABCDEFGHIJK", (14, 32, 16, 22, 10, 30, 45, 45, 12, 35, 35)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "C2"

    ws = wb.create_sheet("Defeitos")
    _header(ws, 1, DEFECT_HEADERS)
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Defeitos Ativos")
    ws["A1"] = "Defeitos ativos (regenerado pelo adaptador — não editar à mão)"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "sem defeitos registados"
    _header(ws, 3, VIEW_HEADERS)

    wb.save(path)
    return path


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    dest = sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "workbook", "plano-demo.xlsx")
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    print(gerar(dest))
