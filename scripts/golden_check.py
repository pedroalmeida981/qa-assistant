"""
Golden test of the Excel adapter (plan, task 1.4; methodology section 15).

Builds the fictional workbook (tests/golden/build_fixture.py) in a temporary
folder, applies tests/golden/batch.json through the adapter with the
configuration adapters/excel/example.yaml, regenerates the active view, saves,
recalculates with LibreOffice and compares with tests/golden/expected.json. Then
applies the same batch again and confirms nothing changes.

    python scripts/golden_check.py [--expected <file>] [--show]
    python scripts/golden_check.py --dir <workspace>/golden/<batch>   (or QA_GOLDEN_DIR)

With --dir (real golden, methodology section 15), after the fictional one: applies the folder's historical
batch to a copy of before.xlsx, with the current configuration of the workspace's project.yaml and the instant
of golden.yaml, and compares every cell with after.xlsx; applying again must change nothing. A batch from before
the translation (plan 8.11) is converted in memory, keeping its identity in the workbook. Real data never leaves
the workspace: the result is only printed.

Exits with 0 if everything matches, 1 if there are differences (one line per
difference), 2 if the environment does not allow the test (e.g. no LibreOffice).
--show prints what was obtained, to review by hand; it never rewrites the expected file.
"""
from __future__ import annotations

import argparse
import contextlib
import io
import json
import os
import shutil
import sys
import tempfile
from datetime import date, datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "tests", "golden"))
sys.path.insert(0, os.path.join(ROOT, "scripts"))

import openpyxl
import yaml
from build_fixture import build
from migrate_workspace import convert_batch, convert_golden, is_old_batch

from adapters.excel import ExcelAdapter, load_batch
from adapters.excel.recalc import find_soffice

GOLDEN = os.path.join(ROOT, "tests", "golden")
CONFIG = os.path.join(ROOT, "adapters", "excel", "example.yaml")
LABEL, REPORTED_BY = "[golden]", "golden"


def _today(obj):
    """Replaces '{today}' with the day's date in every string of the expected file."""
    if isinstance(obj, str):
        return obj.replace("{today}", date.today().isoformat())
    if isinstance(obj, list):
        return [_today(x) for x in obj]
    if isinstance(obj, dict):
        return {k: _today(v) for k, v in obj.items()}
    return obj


def _rows_as_dicts(ws, header_row=1):
    h = [c.value for c in ws[header_row]]
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if row and row[0] not in (None, ""):
            yield dict(zip(h, row))


def _snapshot(path):
    wbf, wbv = openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True)
    return {ws.title: [[(c.value, wbv[ws.title][c.coordinate].value) for c in row] for row in ws.iter_rows()]
            for ws in wbf.worksheets}


def run(workdir: str) -> dict:
    with open(CONFIG, encoding="utf-8") as f:
        cfg = yaml.safe_load(f)
    build(os.path.join(workdir, cfg["file"]))
    batch = load_batch(os.path.join(GOLDEN, "batch.json"))

    ad = ExcelAdapter.from_dict(cfg, base_dir=workdir)
    ad.load()
    r = ad.apply_batch(batch, label=LABEL, reported_by=REPORTED_BY)
    ad.regenerate_active_defects()
    s = ad.save_if_clean()
    rc = ad.recalc() if s["saved_original"] else {"ok": False, "error": "not saved (cross-check)"}

    path = ad.path
    wb = openpyxl.load_workbook(path, data_only=True)
    ob = {
        "apply": {"applied": r["applied"], "new_defects": r["new_defects"]},
        "saved_original": s["saved_original"],
        "cross_check": s["problems"],
        "recalc_ok": rc["ok"],
        "sheets": [[ws.title, ws.sheet_state] for ws in wb.worksheets],
        "cases": {},
        "defects": {},
    }
    if not rc["ok"]:
        ob["recalc_error"] = rc.get("error")

    for sheet in dict.fromkeys(cfg["cases"]["sheets"].values()):
        for row in _rows_as_dicts(wb[sheet]):
            ob["cases"][row["ID"]] = {k: v for k, v in row.items() if k not in ("ID", "Title")}
    for row in _rows_as_dicts(wb[cfg["defects"]["sheet"]]):
        ob["defects"][row["Key"]] = {k: v for k, v in row.items() if k != "Key"}

    v = cfg["active_view"]
    wv = wb[v["sheet"]]
    first = v["header_row"] + 1
    ob["view"] = {
        "A1": wv["A1"].value,
        "A2": wv["A2"].value,
        "header": [c.value for c in wv[v["header_row"]]],
        "rows": [list(row) for row in wv.iter_rows(min_row=first, values_only=True) if row[0] not in (None, "")],
        "rows_bold": all(c.font.bold for row in wv.iter_rows(min_row=first, max_row=wv.max_row) for c in row),
    }

    ws = wb["Summary"]
    h = [c.value for c in ws[1]]
    ob["summary_totals"] = {row[0]: dict(zip(h[1:], row[1:])) for row in ws.iter_rows(min_row=2, values_only=True)}

    # summary (recalculated formulas) vs direct count: technique-excel.md section 5
    vr = ad.check_summary()
    ob["summary_vs_count"] = {"checked": vr["checked"], "problems": vr["problems"]}

    ad_ctl = ExcelAdapter.from_dict(cfg, base_dir=workdir)
    ad_ctl.load()
    ob["batch_control"] = [{k: x[k] for k in ("label", "items", "new_defects")} for x in ad_ctl.applied_batches()]

    # apply the same batch again: nothing may change
    before = _snapshot(path)
    ad2 = ExcelAdapter.from_dict(cfg, base_dir=workdir)
    ad2.load()
    r2 = ad2.apply_batch(batch, label=LABEL, reported_by=REPORTED_BY)
    ad2.regenerate_active_defects()
    s2 = ad2.save_if_clean()
    if s2["saved_original"]:
        ad2.recalc()
    ob["reapply"] = {"applied": r2["applied"], "file_unchanged": _snapshot(path) == before}
    return ob


def compare(expected, obtained, path="") -> list[str]:
    diffs = []
    if isinstance(expected, dict) and isinstance(obtained, dict):
        for k in expected:
            if k.startswith("_"):
                continue
            p = f"{path}.{k}" if path else k
            if k not in obtained:
                diffs.append(f"{p}: expected {expected[k]!r}, missing from the result")
            else:
                diffs += compare(expected[k], obtained[k], p)
        for k in obtained:
            if k not in expected and not k.startswith("_"):
                p = f"{path}.{k}" if path else k
                diffs.append(f"{p}: not expected, obtained {obtained[k]!r}")
    elif isinstance(expected, list) and isinstance(obtained, list) and len(expected) == len(obtained):
        for i, (e, o) in enumerate(zip(expected, obtained)):
            diffs += compare(e, o, f"{path}[{i}]")
    elif expected != obtained:
        diffs.append(f"{path}: expected {expected!r}, obtained {obtained!r}")
    return diffs


# ---------------------------------------------------------------------------
# Real golden (methodology section 15): a workspace's historical batch
# ---------------------------------------------------------------------------


class InvalidGolden(Exception):
    """Real golden folder incomplete or without a project.yaml above (code 2)."""


def _workspace_of(folder: str) -> str:
    """The nearest folder above with a project.yaml (the golden lives inside the workspace)."""
    d = os.path.abspath(folder)
    while True:
        if os.path.isfile(os.path.join(d, "project.yaml")):
            return d
        if os.path.isfile(os.path.join(d, "projeto.yaml")):
            raise InvalidGolden(f"{d} is a workspace from before the translation: convert it with "
                                "scripts/migrate_workspace.py")
        parent = os.path.dirname(d)
        if parent == d:
            raise InvalidGolden(f"no project.yaml in {folder} nor in the folders above (the golden lives in the "
                                "workspace)")
        d = parent


def golden_folders(folder: str) -> list[str]:
    """The folder itself, if it has before.xlsx; otherwise, its subfolders that have it."""
    if os.path.isfile(os.path.join(folder, "before.xlsx")):
        return [folder]
    subs = sorted(os.path.join(folder, d) for d in os.listdir(folder)
                  if os.path.isfile(os.path.join(folder, d, "before.xlsx")))
    if not subs:
        raise InvalidGolden(f"{folder}: no before.xlsx, nor subfolders with before.xlsx")
    return subs


def _cells(path: str) -> dict:
    """{sheet: {cell: (value or formula, calculated value)}}, only cells with content."""
    wbf, wbv = openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True)
    return {ws.title: {c.coordinate: (c.value, wbv[ws.title][c.coordinate].value)
                       for row in ws.iter_rows() for c in row if c.value is not None}
            for ws in wbf.worksheets}


def cell_differences(expected: dict, obtained: dict) -> list[str]:
    diffs = [f"sheet '{s}': expected, missing from the result" for s in expected if s not in obtained]
    diffs += [f"sheet '{s}': not expected" for s in obtained if s not in expected]
    for s, cells_e in expected.items():
        cells_o = obtained.get(s)
        if cells_o is not None:
            for k in sorted(set(cells_e) | set(cells_o)):
                if cells_e.get(k) != cells_o.get(k):
                    diffs.append(f"{s}!{k}: expected {cells_e.get(k)!r}, obtained {cells_o.get(k)!r}")
    return diffs


def _batch_files(folder: str) -> list[str]:
    # the batch keeps its original name, which is its id in the workbook: lote_* before the translation
    return sorted(f for f in os.listdir(folder)
                  if f.startswith(("batch_", "lote_")) and f.endswith(".json"))


def run_real(folder: str) -> list[str]:
    """Applies the historical batch to a copy of before.xlsx through the /checkpoint's apply
    and compares with after.xlsx. Returns the differences (empty = golden OK)."""
    from adapters.excel.apply import apply

    batches = _batch_files(folder)
    missing = [f for f in ("before.xlsx", "after.xlsx", "golden.yaml") if not os.path.isfile(os.path.join(folder, f))]
    if missing or len(batches) != 1:
        raise InvalidGolden(f"{folder}: missing {missing}" if missing else
                            f"{folder}: must have exactly one batch_*.json (has {batches})")
    with open(os.path.join(folder, "golden.yaml"), encoding="utf-8-sig") as f:
        g = convert_golden(yaml.safe_load(f) or {})
    if not g.get("now"):
        raise InvalidGolden(f"{folder}/golden.yaml: 'now' is missing (instant of the original application, "
                            "YYYY-MM-DDTHH:MM:SS)")
    now = g["now"] if isinstance(g["now"], datetime) else datetime.fromisoformat(str(g["now"]))
    with open(os.path.join(_workspace_of(folder), "project.yaml"), encoding="utf-8-sig") as f:
        proj = yaml.safe_load(f) or {}
    if not proj.get("excel"):
        raise InvalidGolden("the workspace's project.yaml has no excel: section")

    diffs = []
    with tempfile.TemporaryDirectory() as tmp:
        # temporary workspace: the current configuration, with the workbook replaced by the copy of before.xlsx
        excel = {**proj["excel"], "file": os.path.basename(proj["excel"]["file"])}
        with open(os.path.join(tmp, "project.yaml"), "w", encoding="utf-8") as f:
            yaml.safe_dump({"excel": excel}, f, allow_unicode=True)
        workbook = os.path.join(tmp, excel["file"])
        shutil.copy(os.path.join(folder, "before.xlsx"), workbook)
        batch = os.path.join(tmp, batches[0])
        with open(os.path.join(folder, batches[0]), encoding="utf-8-sig") as f:
            doc = json.load(f)
        with open(batch, "w", encoding="utf-8") as f:
            json.dump(convert_batch(doc) if is_old_batch(doc) else doc, f, ensure_ascii=False)

        args = {"label": g.get("label"), "reported_by": g.get("reported_by"), "now": now}
        with contextlib.redirect_stdout(io.StringIO()) as out:
            try:
                rc = apply(tmp, batch, **args)
            except Exception as e:  # noqa: BLE001 -- invalid batch, configuration: a difference to investigate
                return [f"apply failed: {e}"]
        if rc not in (0, 5):  # 5 = the workbook's summary is wrong; after.xlsx says whether it was so at the time
            return [f"apply exited with {rc}: {out.getvalue().strip()[-800:]}"]
        diffs += cell_differences(_cells(os.path.join(folder, "after.xlsx")), _cells(workbook))

        before = _cells(workbook)
        with contextlib.redirect_stdout(io.StringIO()) as out:
            apply(tmp, batch, **args)
        if "Nothing applied" not in out.getvalue() or _cells(workbook) != before:
            diffs.append("applying the same batch again changed the workbook")
    return diffs


def main_real(folder: str) -> int:
    try:
        folders = golden_folders(folder)
    except (InvalidGolden, FileNotFoundError) as e:
        print(f"REAL GOLDEN DID NOT RUN: {e}")
        return 2
    failed = False
    for p in folders:
        try:
            diffs = run_real(p)
        except InvalidGolden as e:
            print(f"REAL GOLDEN DID NOT RUN: {e}")
            return 2
        if diffs:
            failed = True
            print(f"REAL GOLDEN FAILED: {p} -- {len(diffs)} difference(s) (investigate each one; methodology "
                  "section 15)")
            for d in diffs[:50]:
                print(f"  - {d}")
            if len(diffs) > 50:
                print(f"  ... {len(diffs) - 50} more")
        else:
            print(f"REAL GOLDEN OK: {p} (cells equal to after.xlsx; applying again has no effect)")
    return 1 if failed else 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Golden test of the Excel adapter")
    ap.add_argument("--expected", default=os.path.join(GOLDEN, "expected.json"))
    ap.add_argument("--show", action="store_true", help="print the result (never writes the expected file)")
    ap.add_argument("--dir", default=os.environ.get("QA_GOLDEN_DIR") or None,
                    help="real golden, IN ADDITION to the fictional one: folder in the workspace with before.xlsx, "
                         "batch_*.json, after.xlsx and golden.yaml (or a folder with several of these); "
                         "default QA_GOLDEN_DIR")
    a = ap.parse_args(argv)

    try:
        find_soffice()
    except FileNotFoundError as e:
        print(f"GOLDEN DID NOT RUN: {e}\n(the Summary totals are only checked with the recalculation)")
        return 2
    rc = fictional_golden(a.expected, a.show)
    # the fictional one always runs (with QA_GOLDEN_DIR in the environment, pre-commit must not lose it)
    return max(rc, main_real(a.dir)) if a.dir else rc


def fictional_golden(expected_path: str, show: bool) -> int:
    with open(expected_path, encoding="utf-8") as f:
        expected = _today(json.load(f))
    with tempfile.TemporaryDirectory() as tmp:
        obtained = run(tmp)

    if show:
        print(json.dumps(obtained, ensure_ascii=False, indent=2, default=str))
    diffs = compare(expected, obtained)
    if diffs:
        print(f"GOLDEN FAILED: {len(diffs)} difference(s) between {os.path.relpath(expected_path, ROOT)} and the "
              "result")
        for d in diffs:
            print(f"  - {d}")
        return 1
    print("GOLDEN OK: the adapter produces exactly the expected result "
          f"({len(expected['cases'])} cases, {len(expected['defects'])} defects, Summary totals, batch applied "
          "again with no effect)")
    return 0


if __name__ == "__main__":
    # on Windows, output read through a pipe (e.g. Claude Code) came in cp1252 and accents arrived garbled
    for _s in (sys.stdout, sys.stderr):
        _s.reconfigure(encoding="utf-8", errors="replace")
    sys.exit(main())
