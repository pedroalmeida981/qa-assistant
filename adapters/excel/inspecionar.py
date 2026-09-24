"""
Inspeciona um workbook e propoe a secao `excel:` do projeto.yaml
(plano, tarefa 1.3; usado pelo /onboarding-projeto).

So LE o workbook -- nunca o grava. Nao adivinha: o que deteta com seguranca
vai preenchido; o resto fica como "CONFIRMAR: ..." com as opcoes encontradas,
para o utilizador escolher. Um "CONFIRMAR" que fique na configuracao faz o
adaptador falhar no load() (coluna inexistente), por isso nao passa despercebido.

    python -m adapters.excel.inspecionar <workbook.xlsx> [--yaml-out proposta.yaml]
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from collections import Counter

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

from .base import _CONTROL_HEADERS, CONTROL_SHEET_DEFAULT

ID_RE = re.compile(r"^[A-Za-z][A-Za-z0-9]*(?:[-_.][A-Za-z0-9]+)+$")
NUM_ID_RE = re.compile(r"^(\D*?)(\d+)(\D*)$")
MAX_HEADER_SCAN = 15
ENUM_MAX_DISTINCT = 15


def _cfm(msg: str, opcoes=None) -> str:
    return f"CONFIRMAR: {msg}" + (f" -- opcoes: {opcoes}" if opcoes else "")


def detect_header_row(ws) -> int | None:
    """Linha (nas primeiras 15) com mais celulas de texto curto, de preferencia
    seguida de dados; numa sheet so com cabecalhos (projeto a comecar), a linha
    com mais texto curto, se tiver pelo menos 3 celulas."""
    best, best_score, fallback, fb_score = None, 1, None, 2
    for r in range(1, min(ws.max_row, MAX_HEADER_SCAN) + 1):
        vals = [c.value for c in ws[r]]
        score = sum(1 for v in vals if isinstance(v, str) and 0 < len(v.strip()) <= 60)
        nxt = [c.value for c in ws[r + 1]] if r < ws.max_row else []
        if score > best_score and any(v not in (None, "") for v in nxt):
            best, best_score = r, score
        if score > fb_score:
            fallback, fb_score = r, score
    return best or fallback


def dv_lists(ws, wb) -> list[tuple[set, list[str]]]:
    """Validacoes de dados do tipo lista: (colunas abrangidas, valores)."""
    out = []
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list" or not dv.formula1:
            continue
        f = dv.formula1.strip()
        if f.startswith('"'):
            vals = [x.strip() for x in f.strip('"').split(",") if x.strip()]
        else:
            ref = f.lstrip("=")
            sheet, _, rng = ref.rpartition("!")
            src = wb[sheet.strip("'")] if sheet else ws
            try:
                c1, r1, c2, r2 = range_boundaries(rng.replace("$", ""))
            except ValueError:
                continue
            vals = [src.cell(r, c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
            vals = [str(v) for v in vals if v not in (None, "")]
        cols = set()
        for rng in str(dv.sqref).split():
            c1, _, c2, _ = range_boundaries(rng)
            cols.update(range(c1, (c2 or c1) + 1))
        out.append((cols, vals))
    return out


def common_token_prefix(ids: list[str]) -> str:
    """Maior prefixo comum, cortado num separador (ex. 'CT-LOJA-')."""
    if not ids:
        return ""
    p = os.path.commonprefix(ids)
    cut = max(p.rfind(ch) for ch in "-_.")
    return p[: cut + 1] if cut >= 0 else ""


def profile_sheet(ws, wb) -> dict:
    h = detect_header_row(ws)
    info = {"sheet": ws.title, "linha_cabecalho": h, "colunas": []}
    if not h:
        return info
    dvs = dv_lists(ws, wb)
    for c in ws[h]:
        if c.value in (None, ""):
            continue
        col = c.column
        vals = [ws.cell(r, col).value for r in range(h + 1, ws.max_row + 1)]
        vals = [v for v in vals if v not in (None, "")]
        svals = [str(v).strip() for v in vals]
        n = len(svals)
        cnt = Counter(svals)
        idr = sum(1 for v in svals if ID_RE.match(v)) / n if n else 0
        lista = next((vs for cols, vs in dvs if col in cols), None)
        info["colunas"].append({
            "nome": c.value, "letra": get_column_letter(col), "n": n,
            "distintos": len(cnt), "id_ratio": idr,
            "unico": n > 0 and len(cnt) == n,
            "valores": cnt.most_common(ENUM_MAX_DISTINCT) if len(cnt) <= ENUM_MAX_DISTINCT else None,
            "lista_validacao": lista,
            "datas": n > 0 and sum(1 for v in vals if hasattr(v, "year") or re.match(r"^\d{4}-\d{2}-\d{2}", str(v))) / n > 0.8,
            "texto_longo": n > 0 and sum(len(v) for v in svals) / n > 40,
            "amostra": svals[:3],
            "todos": svals,
        })
    return info


def classify(profiles: list[dict]) -> dict:
    """Separa sheets de casos, de defeitos e vistas, so pelas relacoes entre
    elas (sem nomes pre-definidos):
    - sheet com ID: coluna com valores tipo ID, todos diferentes;
    - vista: sheet cujos IDs estao todos contidos nos IDs de outra sheet;
    - defeitos: sheet (nao vista) com uma coluna que aponta para IDs de outras sheets;
    - casos: as restantes sheets com ID."""
    res = {"casos": [], "defeitos": [], "vistas": [], "outras": [], "vazias": []}
    com_id = []
    for p in profiles:
        ids = [c for c in p["colunas"] if c["id_ratio"] >= 0.9 and c["unico"] and c["n"] >= 1]
        if p["colunas"] and all(c["n"] == 0 for c in p["colunas"]):
            res["vazias"].append(p)       # so cabecalhos: candidata a defeitos ou vista
        elif ids:
            p["_id"] = ids[0]
            p["_ids"] = set(ids[0]["todos"])
            p["_prefixo"] = common_token_prefix(ids[0]["todos"])
            com_id.append(p)
        else:
            res["outras"].append(p)
    for p in com_id:
        p["_vista_de"] = next((q for q in com_id if q is not p and p["_ids"] <= q["_ids"]), None)
    for p in com_id:
        if p["_vista_de"]:
            continue
        outros = tuple(q["_prefixo"] for q in com_id if q is not p and not q["_vista_de"] and q["_prefixo"])
        # referencia = a maioria dos valores da coluna comeca por um prefixo de ID de outra sheet
        p["_refs"] = [c for c in p["colunas"] if c is not p["_id"] and outros and c["n"]
                      and sum(1 for v in c["todos"] if v.startswith(outros)) / c["n"] >= 0.5]
    for p in com_id:
        if p["_vista_de"]:
            res["vistas"].append(p)
        elif p["_refs"]:
            # formato seguido por quase todos os IDs; os que nao o seguem ficam como excecoes
            todos = p["_id"]["todos"]
            forms = Counter((m.group(1), m.group(3)) for m in map(NUM_ID_RE.match, todos) if m)
            (pre, pos), n = forms.most_common(1)[0] if forms else (("", ""), 0)
            p["_excecoes_id"] = [v for v in todos if not (NUM_ID_RE.match(v) and
                                 (NUM_ID_RE.match(v).group(1), NUM_ID_RE.match(v).group(3)) == (pre, pos))]
            if todos and n / len(todos) >= 0.95:
                p["_formato_id"] = f"{pre}{{n}}{pos}"
            else:
                p["_formato_id"] = _cfm("formato do ID de defeito com {n}", p["_id"]["amostra"])
            res["defeitos"].append(p)
        else:
            res["casos"].append(p)
    return res


def _enum_cols(p, excluir=()):
    return [c for c in p["colunas"] if c["nome"] not in excluir and
            (c["lista_validacao"] or (c["valores"] and 1 < c["distintos"] <= ENUM_MAX_DISTINCT
                                      and c["n"] >= 3 and not c["texto_longo"] and not c["datas"]
                                      and (c["n"] < 20 or c["distintos"] <= c["n"] / 2)))]


def _estado_proposta(p, excluir):
    enums = _enum_cols(p, excluir)
    com_lista = [c for c in enums if c["lista_validacao"]]
    if len(com_lista) == 1:
        c = com_lista[0]
        return c["nome"], c["lista_validacao"], "detetado pela validacao de dados (lista)"
    return (_cfm("coluna de estado", [c["nome"] for c in enums]),
            None, "varias colunas candidatas" if enums else "sem candidatas")


def _match_col(c, base) -> str:
    """Coluna da sheet de defeitos cujos valores contem os valores desta coluna da vista."""
    if not c["n"]:
        return _cfm("coluna vazia na vista -- sem origem visivel na sheet de defeitos; "
                    "apagar esta linha ou indicar a coluna")
    vals = set(c["todos"])
    hits = [b["nome"] for b in base["colunas"] if b["n"] and len(vals & set(b["todos"])) / len(vals) >= 0.9]
    if len(hits) == 1:
        return hits[0]
    return _cfm(f"coluna de '{base['sheet']}' correspondente", hits or [b["nome"] for b in base["colunas"]])


def propose(path: str) -> tuple[dict, list[str]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    # a sheet de controlo do proprio adaptador nao faz parte do layout do projeto
    controlo = [ws.title for ws in wb.worksheets if ws.title == CONTROL_SHEET_DEFAULT
                or [c.value for c in ws[1]][:len(_CONTROL_HEADERS)] == _CONTROL_HEADERS]
    profiles = [profile_sheet(ws, wb) for ws in wb.worksheets if ws.title not in controlo]
    cls = classify(profiles)
    notas = [f"sheet '{t}': controlo de lotes do adaptador -- ignorada" for t in controlo]

    # ---- casos
    sheets, estados_caso, id_cols, est_cols = {}, [], set(), set()
    for p in cls["casos"]:
        idc = p["_id"]
        pre = p["_prefixo"]
        sheets[pre or _cfm(f"prefixo dos IDs de '{p['sheet']}'")] = p["sheet"]
        id_cols.add(idc["nome"])
        nome, estados, como = _estado_proposta(p, {idc["nome"]})
        est_cols.add(nome)
        estados_caso += [e for e in (estados or []) if e not in estados_caso]
        notas.append(f"casos '{p['sheet']}': cabecalho na linha {p['linha_cabecalho']}, ID '{idc['nome']}' "
                     f"(prefixo '{pre}'), estado '{nome}' ({como})")
    livres = sorted({c["nome"] for p in cls["casos"] for c in p["colunas"]
                     if c["texto_longo"] or c["n"] == 0})
    linhas_cab = {p["linha_cabecalho"] for p in cls["casos"]}

    casos = {
        "sheets": sheets,
        "linha_cabecalho": linhas_cab.pop() if len(linhas_cab) == 1 else _cfm("linhas diferentes", sorted(linhas_cab)),
        "colunas": {
            "id": id_cols.pop() if len(id_cols) == 1 else _cfm("coluna de ID", sorted(id_cols)),
            "estado": est_cols.pop() if len(est_cols) == 1 else _cfm("coluna de estado", sorted(est_cols)),
            "resultado_obtido": _cfm("coluna de resultado obtido (opcional)", livres),
            "observacoes": _cfm("coluna de observacoes (opcional)", livres),
        },
        "estados": estados_caso or [_cfm("estados validos de caso (valores da coluna de estado)")],
    }

    # ---- defeitos
    defeitos = {"sheet": _cfm("sheet de defeitos", [p["sheet"] for p in cls["defeitos"]])}
    if len(cls["defeitos"]) == 1:
        p = cls["defeitos"][0]
        idc = p["_id"]
        refs = [c["nome"] for c in p["_refs"]]
        caso_col = refs[0] if len(refs) == 1 else _cfm("coluna do caso relacionado", refs)
        nome, estados, como = _estado_proposta(p, {idc["nome"]})
        enums = [c["nome"] for c in _enum_cols(p, {idc["nome"], nome})]
        datas = [c["nome"] for c in p["colunas"] if c["datas"]]
        defeitos = {
            "sheet": p["sheet"],
            "linha_cabecalho": p["linha_cabecalho"],
            "formato_id": p["_formato_id"],
            "colunas": {
                "id": idc["nome"], "estado": nome,
                "caso": caso_col,
                "severidade": _cfm("coluna de severidade (opcional)", enums),
                "prioridade": _cfm("coluna de prioridade (opcional)", enums),
                "reportado_por": _cfm("coluna de quem reportou (opcional)"),
                "data_report": _cfm("coluna da data de report (opcional)", datas),
            },
            "campos_lote": {"titulo": _cfm("coluna do titulo"), "descricao": _cfm("coluna da descricao")},
            "valores_omissao": {},
            "estados": estados or [_cfm("estados validos de defeito (valores da coluna de estado)")],
            "estados_fechados": [_cfm("estados que contam como resolvidos", estados)],
            "estado_inicial": _cfm("estado de um defeito novo", estados),
            "prioridades": [_cfm("ordem das prioridades, da mais alta para a mais baixa")],
        }
        notas.append(f"defeitos '{p['sheet']}': cabecalho na linha {p['linha_cabecalho']}, ID '{idc['nome']}' "
                     f"(formato '{p['_formato_id']}'), estado '{nome}' ({como}), caso relacionado '{caso_col}'")
        if p["_excecoes_id"]:
            notas.append(f"   IDs que nao seguem o formato (ignorados na numeracao): {p['_excecoes_id'][:10]}")
        for c in _enum_cols(p, {idc["nome"]}):
            vs = c["lista_validacao"] or [v for v, _ in c["valores"]]
            notas.append(f"   coluna '{c['nome']}': {vs}")

    # ---- sem sheet de defeitos com dados: propor a partir das sheets so com cabecalhos
    if not cls["defeitos"] and cls["vazias"]:
        nomes = [p["sheet"] for p in cls["vazias"]]
        cab = {p["sheet"]: [c["nome"] for c in p["colunas"]] for p in cls["vazias"]}
        opc = lambda what: _cfm(what, cab)
        defeitos = {
            "sheet": _cfm("sheet de defeitos (sem defeitos ainda)", nomes),
            "linha_cabecalho": _cfm("linha dos cabecalhos", {p["sheet"]: p["linha_cabecalho"] for p in cls["vazias"]}),
            "formato_id": _cfm("formato do ID de defeito com {n}, ex. BUG-{n}"),
            "colunas": {k: opc(f"coluna de {k}") for k in
                        ("id", "estado", "caso", "severidade", "prioridade", "reportado_por", "data_report")},
            "campos_lote": {"titulo": opc("coluna do titulo"), "descricao": opc("coluna da descricao")},
            "valores_omissao": {},
            "estados": [_cfm("estados validos de defeito")],
            "estados_fechados": [_cfm("estados que contam como resolvidos")],
            "estado_inicial": _cfm("estado de um defeito novo"),
            "prioridades": [_cfm("ordem das prioridades, da mais alta para a mais baixa")],
        }
        notas.append(f"sem defeitos registados: sheets so com cabecalhos {cab} -- "
                     f"uma deve ser a de defeitos e, se existir, outra a vista de ativos")

    # ---- vista de ativos: outra sheet cujo ID e do mesmo formato dos defeitos
    vista = None
    if not cls["defeitos"] and len(cls["vazias"]) > 1:
        vista = {"sheet": _cfm("sheet da vista de ativos (opcional)", [p["sheet"] for p in cls["vazias"]]),
                 "linha_cabecalho": _cfm("linha dos cabecalhos da vista"),
                 "colunas": _cfm("cabecalho da vista -> coluna dos defeitos (so os que tiverem nome diferente)"),
                 "ordenar_por": "", "resumo": None}
    for p in cls["vistas"]:
        base = p["_vista_de"]
        base_cols = [x["nome"] for x in base["colunas"]]
        vista = {"sheet": p["sheet"], "linha_cabecalho": p["linha_cabecalho"],
                 "colunas": {c["nome"]: (base["_id"]["nome"] if c is p["_id"] else _match_col(c, base))
                             for c in p["colunas"] if c["nome"] not in base_cols or c is p["_id"]},
                 "ordenar_por": "", "resumo": None}
        acima = [(c.coordinate, str(c.value)[:60]) for row in wb[p["sheet"]].iter_rows(max_row=p["linha_cabecalho"] - 1)
                 for c in row if c.value not in (None, "")]
        notas.append(f"vista de ativos candidata: '{p['sheet']}' (IDs contidos em '{base['sheet']}'; "
                     f"cabecalho na linha {p['linha_cabecalho']}); "
                     f"celulas acima do cabecalho: {acima}")
    if len(cls["vistas"]) > 1:
        notas.append(f"ATENCAO: {len(cls['vistas'])} vistas candidatas -- so a ultima entrou na proposta: "
                     f"{[p['sheet'] for p in cls['vistas']]}")
    for p in cls["outras"]:
        notas.append(f"sheet '{p['sheet']}': sem coluna de ID -- nao entra na configuracao")

    cfg = {
        "ficheiro": _cfm("caminho do workbook relativo ao projeto.yaml", [os.path.basename(path)]),
        "casos": casos, "defeitos": defeitos, "vista_ativos": vista,
        "recalc": {"mecanismo": "libreoffice", "timeout": 60},
    }
    return cfg, notas


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("workbook")
    ap.add_argument("--yaml-out", help="gravar a proposta neste ficheiro (o workbook nunca e gravado)")
    a = ap.parse_args(argv)
    import yaml

    cfg, notas = propose(a.workbook)
    print("# Evidencia encontrada:")
    for n in notas:
        print("#  -", n)
    out = yaml.safe_dump({"excel": cfg}, allow_unicode=True, sort_keys=False, width=200)
    n_cfm = out.count("CONFIRMAR")
    print(f"# {n_cfm} valor(es) por CONFIRMAR com o utilizador.\n")
    print(out)
    if a.yaml_out:
        with open(a.yaml_out, "w", encoding="utf-8") as f:
            f.write(out)


if __name__ == "__main__":
    sys.exit(main())
