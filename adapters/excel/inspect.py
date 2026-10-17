"""
Inspects a workbook and proposes the `excel:` section of project.yaml
(plan, task 1.3; used by /project-onboarding).

Only READS the workbook -- never saves it. Does not guess: what it detects
safely is filled in; the rest stays as "CONFIRM: ..." with the options found,
for the user to choose. A "CONFIRM" left in the configuration makes the
adapter fail in load() (column does not exist), so it does not go unnoticed.

    python -m adapters.excel.inspect <workbook.xlsx> [--yaml-out proposal.yaml]
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from collections import Counter

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

from .base import _CONTROL_HEADERS, _LEGACY_CONTROL_HEADERS, CONTROL_SHEET_DEFAULT

ID_RE = re.compile(r"^[A-Za-z][A-Za-z0-9]*(?:[-_.][A-Za-z0-9]+)+$")
NUM_ID_RE = re.compile(r"^(\D*?)(\d+)(\D*)$")
MAX_HEADER_SCAN = 15
ENUM_MAX_DISTINCT = 15
# the adapter's own control sheet, also as written before the translation (plan 8.11)
CONTROL_SHEETS = (CONTROL_SHEET_DEFAULT, "_qa_lotes")


def _cfm(msg: str, options=None) -> str:
    return f"CONFIRM: {msg}" + (f" -- options: {options}" if options else "")


def case_fields(profiles: list[dict], taken: set) -> dict:
    """cases.fields: each field of the case model whose column is recognised by its header (adapters.case_headers,
    in the common languages) in the case sheets. A header that differs between sheets stays CONFIRM. The ID and the
    status columns are not fields (cases.columns)."""
    from adapters.case_headers import field_of
    found: dict[str, set] = {}
    for p in profiles:
        for c in p["columns"]:
            field = field_of(c["name"])
            if field and field != "id" and c["name"] not in taken:
                found.setdefault(field, set()).add(c["name"])
    order = ("title", "expected", "steps", "preconditions", "role", "priority", "technique", "platform", "area",
             "module")
    return {f: (next(iter(found[f])) if len(found[f]) == 1 else _cfm(f"{f} column", sorted(found[f])))
            for f in order if f in found}


def detect_header_row(ws) -> int | None:
    """Row (among the first 15) with the most short text cells, preferably
    followed by data; in a sheet with headers only (a project just starting),
    the row with the most short text, if it has at least 3 cells."""
    best, best_score, fallback, fb_score = None, 1, None, 2
    for r in range(1, min(ws.max_row, MAX_HEADER_SCAN) + 1):
        vals = [c.value for c in ws[r]]
        score = sum(1 for v in vals if isinstance(v, str) and 0 < len(v.strip()) <= 60)
        nxt = [c.value for c in ws[r + 1]] if r < ws.max_row else []
        if score > best_score and any(v not in (None, "") for v in nxt):
            best, best_score = r, score
        if score > fb_score:
            fallback, fb_score = r, score
    return best or fallback


def dv_lists(ws, wb) -> list[tuple[set, list[str]]]:
    """Data validations of type list: (columns covered, values)."""
    out = []
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list" or not dv.formula1:
            continue
        f = dv.formula1.strip()
        if f.startswith('"'):
            vals = [x.strip() for x in f.strip('"').split(",") if x.strip()]
        else:
            ref = f.lstrip("=")
            sheet, _, rng = ref.rpartition("!")
            src = wb[sheet.strip("'")] if sheet else ws
            try:
                c1, r1, c2, r2 = range_boundaries(rng.replace("$", ""))
            except ValueError:
                continue
            vals = [src.cell(r, c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
            vals = [str(v) for v in vals if v not in (None, "")]
        cols = set()
        for rng in str(dv.sqref).split():
            c1, _, c2, _ = range_boundaries(rng)
            cols.update(range(c1, (c2 or c1) + 1))
        out.append((cols, vals))
    return out


def common_token_prefix(ids: list[str]) -> str:
    """Longest common prefix, cut at a separator (e.g. 'TC-SHOP-')."""
    if not ids:
        return ""
    p = os.path.commonprefix(ids)
    cut = max(p.rfind(ch) for ch in "-_.")
    return p[: cut + 1] if cut >= 0 else ""


def profile_sheet(ws, wb) -> dict:
    h = detect_header_row(ws)
    info = {"sheet": ws.title, "header_row": h, "columns": []}
    if not h:
        return info
    dvs = dv_lists(ws, wb)
    for c in ws[h]:
        if c.value in (None, ""):
            continue
        col = c.column
        vals = [ws.cell(r, col).value for r in range(h + 1, ws.max_row + 1)]
        vals = [v for v in vals if v not in (None, "")]
        svals = [str(v).strip() for v in vals]
        n = len(svals)
        cnt = Counter(svals)
        idr = sum(1 for v in svals if ID_RE.match(v)) / n if n else 0
        validation = next((vs for cols, vs in dvs if col in cols), None)
        info["columns"].append({
            "name": c.value, "letter": get_column_letter(col), "n": n,
            "distinct": len(cnt), "id_ratio": idr,
            "unique": n > 0 and len(cnt) == n,
            "values": cnt.most_common(ENUM_MAX_DISTINCT) if len(cnt) <= ENUM_MAX_DISTINCT else None,
            "validation_list": validation,
            "dates": n > 0 and sum(1 for v in vals if hasattr(v, "year") or re.match(r"^\d{4}-\d{2}-\d{2}", str(v))) / n > 0.8,
            "long_text": n > 0 and sum(len(v) for v in svals) / n > 40,
            "sample": svals[:3],
            "all": svals,
        })
    return info


def _is_view_of(p: dict, q: dict) -> bool:
    """p is a view of q: the IDs of p are all in q. With the SAME IDs (e.g. every
    defect open, the view equal to the register), the view is the one with rows
    above the header (title, summary); on a tie, the one with fewer columns.
    With no tie-break, neither is a view (left to CONFIRM)."""
    if not p["_ids"] <= q["_ids"]:
        return False
    if p["_ids"] < q["_ids"]:
        return True
    if p["header_row"] != q["header_row"]:
        return p["header_row"] > q["header_row"]
    return len(p["columns"]) < len(q["columns"])


def classify(profiles: list[dict]) -> dict:
    """Separates case sheets, defect sheets and views, only by the relations
    between them (no predefined names):
    - sheet with an ID: a column with ID-like values, all different;
    - view: a sheet whose IDs are all contained in another sheet's IDs;
    - defects: a sheet (not a view) with a column pointing to other sheets' IDs;
    - cases: the other sheets with an ID."""
    res = {"cases": [], "defects": [], "views": [], "other": [], "empty": []}
    with_id = []
    for p in profiles:
        ids = [c for c in p["columns"] if c["id_ratio"] >= 0.9 and c["unique"] and c["n"] >= 1]
        if p["columns"] and all(c["n"] == 0 for c in p["columns"]):
            res["empty"].append(p)       # headers only: candidate for defects or view
        elif ids:
            p["_id"] = ids[0]
            p["_ids"] = set(ids[0]["all"])
            p["_prefix"] = common_token_prefix(ids[0]["all"])
            with_id.append(p)
        else:
            res["other"].append(p)
    for p in with_id:
        p["_view_of"] = next((q for q in with_id if q is not p and _is_view_of(p, q)), None)
    for p in with_id:
        if p["_view_of"]:
            continue
        others = tuple(q["_prefix"] for q in with_id if q is not p and not q["_view_of"] and q["_prefix"])
        # reference = most of the column's values start with another sheet's ID prefix
        p["_refs"] = [c for c in p["columns"] if c is not p["_id"] and others and c["n"]
                      and sum(1 for v in c["all"] if v.startswith(others)) / c["n"] >= 0.5]
    for p in with_id:
        if p["_view_of"]:
            res["views"].append(p)
        elif p["_refs"]:
            # format followed by almost every ID; those that do not follow it are exceptions
            every = p["_id"]["all"]
            forms = Counter((m.group(1), m.group(3)) for m in map(NUM_ID_RE.match, every) if m)
            (pre, post), n = forms.most_common(1)[0] if forms else (("", ""), 0)
            p["_id_exceptions"] = [v for v in every if not (NUM_ID_RE.match(v) and
                                   (NUM_ID_RE.match(v).group(1), NUM_ID_RE.match(v).group(3)) == (pre, post))]
            if every and n / len(every) >= 0.95:
                p["_id_format"] = f"{pre}{{n}}{post}"
            else:
                p["_id_format"] = _cfm("defect ID format with {n}", p["_id"]["sample"])
            res["defects"].append(p)
        else:
            res["cases"].append(p)
    return res


def _enum_cols(p, exclude=()):
    return [c for c in p["columns"] if c["name"] not in exclude and
            (c["validation_list"] or (c["values"] and 1 < c["distinct"] <= ENUM_MAX_DISTINCT
                                      and c["n"] >= 3 and not c["long_text"] and not c["dates"]
                                      and (c["n"] < 20 or c["distinct"] <= c["n"] / 2)))]


def _status_proposal(p, exclude):
    enums = _enum_cols(p, exclude)
    with_list = [c for c in enums if c["validation_list"]]
    if len(with_list) == 1:
        c = with_list[0]
        return c["name"], c["validation_list"], "detected by the data validation (list)"
    # a single value (e.g. every defect still "Open") can also be the status column
    single = [c for c in p["columns"] if c["name"] not in exclude and c["distinct"] == 1
              and not c["long_text"] and not c["dates"] and c not in enums]
    options = [c["name"] for c in enums + single]
    return (_cfm("status column", options), None, "several candidate columns" if options else "no candidates")


def _match_col(c, base) -> str:
    """Column of the defects sheet whose values contain the values of this view column."""
    if not c["n"]:
        return _cfm("empty column in the view -- no visible source in the defects sheet; "
                    "delete this line or name the column")
    vals = set(c["all"])
    hits = [b["name"] for b in base["columns"] if b["n"] and len(vals & set(b["all"])) / len(vals) >= 0.9]
    if len(hits) == 1:
        return hits[0]
    return _cfm(f"matching column of '{base['sheet']}'", hits or [b["name"] for b in base["columns"]])


def _is_control_sheet(ws) -> bool:
    head = [c.value for c in ws[1]][:len(_CONTROL_HEADERS)]
    return ws.title in CONTROL_SHEETS or [_LEGACY_CONTROL_HEADERS.get(h, h) for h in head] == _CONTROL_HEADERS


def propose(path: str) -> tuple[dict, list[str]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    # the adapter's own control sheet is not part of the project's layout
    control = [ws.title for ws in wb.worksheets if _is_control_sheet(ws)]
    profiles = [profile_sheet(ws, wb) for ws in wb.worksheets if ws.title not in control]
    cls = classify(profiles)
    notes = [f"sheet '{t}': the adapter's batch control -- ignored" for t in control]

    # ---- cases
    sheets, case_statuses, id_cols, st_cols = {}, [], set(), set()
    for p in cls["cases"]:
        idc = p["_id"]
        pre = p["_prefix"]
        sheets[pre or _cfm(f"ID prefix of '{p['sheet']}'")] = p["sheet"]
        id_cols.add(idc["name"])
        name, statuses, how = _status_proposal(p, {idc["name"]})
        st_cols.add(name)
        case_statuses += [e for e in (statuses or []) if e not in case_statuses]
        notes.append(f"cases '{p['sheet']}': header on row {p['header_row']}, ID '{idc['name']}' "
                     f"(prefix '{pre}'), status '{name}' ({how})")
    free = sorted({c["name"] for p in cls["cases"] for c in p["columns"]
                   if c["long_text"] or c["n"] == 0})
    header_rows = {p["header_row"] for p in cls["cases"]}

    cases = {
        "sheets": sheets,
        "header_row": header_rows.pop() if len(header_rows) == 1 else _cfm("different rows", sorted(header_rows)),
        "columns": {
            "id": id_cols.pop() if len(id_cols) == 1 else _cfm("ID column", sorted(id_cols)),
            "status": st_cols.pop() if len(st_cols) == 1 else _cfm("status column", sorted(st_cols)),
            "actual_result": _cfm("actual result column (optional)", free),
            "notes": _cfm("notes column (optional)", free),
        },
        "statuses": case_statuses or [_cfm("valid case statuses (values of the status column)")],
    }
    fields = case_fields(cls["cases"], set(cases["columns"].values()))
    if fields:
        cases["fields"] = fields

    # ---- defects
    defects = {"sheet": _cfm("defects sheet", [p["sheet"] for p in cls["defects"]])}
    if len(cls["defects"]) == 1:
        p = cls["defects"][0]
        idc = p["_id"]
        refs = [c["name"] for c in p["_refs"]]
        case_col = refs[0] if len(refs) == 1 else _cfm("related case column", refs)
        name, statuses, how = _status_proposal(p, {idc["name"], *refs})
        enums = [c["name"] for c in _enum_cols(p, {idc["name"], name, *refs})]
        dates = [c["name"] for c in p["columns"] if c["dates"]]
        defects = {
            "sheet": p["sheet"],
            "header_row": p["header_row"],
            "id_format": p["_id_format"],
            "columns": {
                "id": idc["name"], "status": name,
                "case": case_col,
                "severity": _cfm("severity column (optional)", enums),
                "priority": _cfm("priority column (optional)", enums),
                "reported_by": _cfm("reporter column (optional)"),
                "report_date": _cfm("report date column (optional)", dates),
            },
            "batch_fields": {"title": _cfm("title column"), "description": _cfm("description column")},
            "defaults": {},
            "statuses": statuses or [_cfm("valid defect statuses (values of the status column)")],
            "closed_statuses": [_cfm("statuses that count as resolved", statuses)],
            "initial_status": _cfm("status of a new defect", statuses),
            "priorities": [_cfm("order of the priorities, highest first")],
        }
        notes.append(f"defects '{p['sheet']}': header on row {p['header_row']}, ID '{idc['name']}' "
                     f"(format '{p['_id_format']}'), status '{name}' ({how}), related case '{case_col}'")
        if p["_id_exceptions"]:
            notes.append(f"   IDs that do not follow the format (ignored in the numbering): {p['_id_exceptions'][:10]}")
        for c in _enum_cols(p, {idc["name"]}):
            vs = c["validation_list"] or [v for v, _ in c["values"]]
            notes.append(f"   column '{c['name']}': {vs}")

    # ---- no defects sheet with data: propose from the sheets with headers only
    if not cls["defects"] and cls["empty"]:
        names = [p["sheet"] for p in cls["empty"]]
        head = {p["sheet"]: [c["name"] for c in p["columns"]] for p in cls["empty"]}
        opt = lambda what: _cfm(what, head)
        defects = {
            "sheet": _cfm("defects sheet (no defects yet)", names),
            "header_row": _cfm("header row", {p["sheet"]: p["header_row"] for p in cls["empty"]}),
            "id_format": _cfm("defect ID format with {n}, e.g. BUG-{n}"),
            "columns": {k: opt(f"{k} column") for k in
                        ("id", "status", "case", "severity", "priority", "reported_by", "report_date")},
            "batch_fields": {"title": opt("title column"), "description": opt("description column")},
            "defaults": {},
            "statuses": [_cfm("valid defect statuses")],
            "closed_statuses": [_cfm("statuses that count as resolved")],
            "initial_status": _cfm("status of a new defect"),
            "priorities": [_cfm("order of the priorities, highest first")],
        }
        notes.append(f"no defects recorded: sheets with headers only {head} -- "
                     f"one should be the defects sheet and, if it exists, another the active view")

    # ---- active view: another sheet whose ID has the defects' format
    view = None
    if not cls["defects"] and len(cls["empty"]) > 1:
        view = {"sheet": _cfm("active view sheet (optional)", [p["sheet"] for p in cls["empty"]]),
                "header_row": _cfm("header row of the view"),
                "columns": _cfm("view header -> defects column (only those with a different name)"),
                "sort_by": "", "summary": None}
    for p in cls["views"]:
        base = p["_view_of"]
        base_cols = [x["name"] for x in base["columns"]]
        view = {"sheet": p["sheet"], "header_row": p["header_row"],
                "columns": {c["name"]: (base["_id"]["name"] if c is p["_id"] else _match_col(c, base))
                            for c in p["columns"] if c["name"] not in base_cols or c is p["_id"]},
                "sort_by": "", "summary": None}
        above = [(c.coordinate, str(c.value)[:60]) for row in wb[p["sheet"]].iter_rows(max_row=p["header_row"] - 1)
                 for c in row if c.value not in (None, "")]
        notes.append(f"candidate active view: '{p['sheet']}' (IDs contained in '{base['sheet']}'; "
                     f"header on row {p['header_row']}); "
                     f"cells above the header: {above}")
    if len(cls["views"]) > 1:
        notes.append(f"WARNING: {len(cls['views'])} candidate views -- only the last one is in the proposal: "
                     f"{[p['sheet'] for p in cls['views']]}")
    for p in cls["other"]:
        notes.append(f"sheet '{p['sheet']}': no ID column -- not part of the configuration")

    cfg = {
        "file": _cfm("workbook path relative to project.yaml", [os.path.basename(path)]),
        "cases": cases, "defects": defects, "active_view": view,
        "recalc": {"mechanism": "libreoffice", "timeout": 60},
    }
    return cfg, notes


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("workbook")
    ap.add_argument("--yaml-out", help="save the proposal to this file (the workbook is never saved)")
    a = ap.parse_args(argv)
    import yaml

    cfg, notes = propose(a.workbook)
    print("# Evidence found:")
    for n in notes:
        print("#  -", n)
    out = yaml.safe_dump({"excel": cfg}, allow_unicode=True, sort_keys=False, width=200)
    n_cfm = out.count("CONFIRM")
    print(f"# {n_cfm} value(s) to CONFIRM with the user.\n")
    print(out)
    if a.yaml_out:
        with open(a.yaml_out, "w", encoding="utf-8") as f:
            f.write(out)


if __name__ == "__main__":
    # on Windows, output read through a pipe (e.g. Claude Code) came in cp1252 and accents arrived garbled
    for _s in (sys.stdout, sys.stderr):
        _s.reconfigure(encoding="utf-8", errors="replace")
    sys.exit(main())
