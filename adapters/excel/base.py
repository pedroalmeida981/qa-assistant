"""
Universal Excel adapter (technique-excel.md section 1).

Applies a batch of results (run log, methodology section 7.1) to an Excel
workbook used as the source of truth for test cases and defects:

    load (keeping formulas) -> apply batch -> regenerate the active defects
    view -> cross-check -> save only if clean (else .REVIEWED.xlsx)
    -> recalculate safely

No sheet, column or status name lives in this code: the whole layout comes
from the configuration (see example.yaml next to this file; fields documented
in from_dict()).

Typical use:

    from adapters.excel import ExcelAdapter, load_batch

    ad = ExcelAdapter.from_yaml("project.yaml", section="excel")
    ad.load()
    res = ad.apply_batch(load_batch("runs/batch_x.json"), label="[2026-09-23]")
    ad.regenerate_active_defects()
    out = ad.save_if_clean()
    if out["saved_original"]:
        ad.recalc()

Batch format: {"results": [item, ...]}, each item
    {"tc": <case id>, "status": <status>, "actual_result": <text>,
     "new_defect": null | {<key>: <value>, ...}}
The keys of "new_defect" are mapped to columns by the configuration
(defects.batch_fields); "related_tc" (by default the case itself),
"severity" and "priority" (by default the severity) have their own role.
"bug" (optional) links an existing defect to the case (a note only).
Other keys (ts, note, ...) are ignored. An invalid batch is refused whole
before anything is written (validate_batch, technique-excel.md section 4).
"""
from __future__ import annotations

import copy
import hashlib
import json
import os
import re
from collections import Counter
from datetime import datetime

import openpyxl
from openpyxl.formula.translate import Translator

from .recalc import recalc_safe

CONTROL_SHEET_DEFAULT = "_qa_batches"
_CONTROL_HEADERS = ["hash", "applied_at", "label", "items", "new_defects", "batch_id"]
# control sheets written before the translation (plan 8.11) keep their headers: they are workbook data
_LEGACY_CONTROL_HEADERS = {"aplicado_em": "applied_at", "etiqueta": "label", "itens": "items",
                           "defeitos_novos": "new_defects", "lote_id": "batch_id"}
REPORTED_BY_DEFAULT = "Automated QA"
# fields of the case model (methodology section 2) that add-cases can write (technique-excel.md section 4.1)
CASE_MODEL_FIELDS = ("title", "area", "feature", "preconditions", "steps", "expected", "role", "platform",
                     "priority", "technique", "origin")


class ConfigError(ValueError):
    pass


class InvalidBatch(ValueError):
    """Batch refused before anything is written; .problems holds the list."""

    def __init__(self, problems: list[str]):
        self.problems = problems
        super().__init__("invalid batch:\n - " + "\n - ".join(problems))


def read_batch(path: str) -> dict:
    """The whole batch file: {"results": [...], "translated_from": <hash>?, ...}."""
    with open(path, encoding="utf-8-sig") as f:  # -sig: accepts files saved with a BOM
        data = json.load(f)
    if not isinstance(data, dict) or not isinstance(data.get("results"), list):
        hint = (" (a batch from before the translation: convert the workspace with scripts/migrate_workspace.py)"
                if isinstance(data, dict) and "resultados" in data else "")
        raise InvalidBatch([f"{path}: expected {{\"results\": [...]}} (technique-excel.md section 4){hint}"])
    return data


def load_batch(path: str) -> list:
    return read_batch(path)["results"]


def batch_hash(batch: list) -> str:
    """Fingerprint of the batch's content (independent of key order)."""
    canon = json.dumps(batch, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(canon.encode("utf-8")).hexdigest()


def _req(d: dict, key: str, where: str):
    if key not in d or d[key] in (None, "", [], {}):
        raise ConfigError(f"configuration: '{where}{key}' is missing")
    return d[key]


class ExcelAdapter:
    # ------------------------------------------------------------------
    # Configuration
    # ------------------------------------------------------------------

    def __init__(self, config: dict, base_dir: str | None = None):
        self.cfg = self._validate(config)
        self.base_dir = base_dir or os.getcwd()
        self.wb: openpyxl.Workbook | None = None
        # instant used in the dates written (label, report date, view summary, applied_at);
        # None = the clock. Fixed, the same sequence of batches always gives the same cells
        # (technique-excel.md section 1, "Date of the application")
        self.now: datetime | None = None

    def _now(self) -> datetime:
        return self.now or datetime.now()

    @classmethod
    def from_dict(cls, config: dict, base_dir: str | None = None) -> ExcelAdapter:
        """Fields (required unless stated):

        file: path of the workbook (relative to base_dir)
        cases:
          sheets: {id_prefix: sheet_name}
          header_row: int (optional, 1)
          columns: {id, status, actual_result (opt.), notes (opt.)}
                   (the optional ones are used only in the sheets that have them)
          statuses: [valid case statuses]
          fields: {model_field: column} (opt.; add-cases, title required there)
          new_status: status of a case added by add-cases (opt., the 1st status)
        defects:
          sheet, header_row (opt., 1)
          id_format: e.g. "BUG-{n}"
          columns: {id, status, case (opt.), severity (opt.), priority (opt.),
                    reported_by (opt.), report_date (opt.)}
          batch_fields: {defect_key_in_the_batch: column} (opt.)
          defaults: {defect_key_in_the_batch: value} (opt.)
          reported_by: value of the reporter column (opt., "Automated QA")
          statuses, closed_statuses, initial_status
          priorities: [order, highest first] (opt.)
        active_view (opt.):
          sheet, header_row (opt., 1)
          columns: {view_header: defects_column} (unmapped headers use the
                   column with the same name)
          sort_by: defects column used as the 2nd criterion (opt.)
          summary: {cell, template} (opt.; template with {open}, {closed},
                   {severity[<value>]}, {date:%d.%m.%Y})
        notes (opt.):
          result: template with {result} (used when there is no actual_result column)
          new_defect: template with {id}
          existing_defect: template with {id} (item with "bug")
        batch_control (opt.): {sheet} -- hidden sheet with the batches applied
        recalc (opt.): {mechanism: libreoffice | excel-com, timeout: int}
        """
        return cls(config, base_dir)

    @classmethod
    def from_yaml(cls, path: str, section: str | None = None) -> ExcelAdapter:
        import yaml

        with open(path, encoding="utf-8-sig") as f:
            data = yaml.safe_load(f) or {}
        if section:
            if section not in data:
                raise ConfigError(f"{path}: section '{section}' is missing")
            data = data[section]
        return cls(data, base_dir=os.path.dirname(os.path.abspath(path)))

    @classmethod
    def from_project(cls, workspace: str = ".") -> ExcelAdapter:
        """Reads the `excel:` section of the workspace's project.yaml (plan
        task 1.3). The workbook path is relative to the workspace."""
        path = os.path.join(workspace, "project.yaml")
        if not os.path.exists(path):
            old = os.path.join(workspace, "projeto.yaml")
            hint = (" (a workspace from before the translation: convert it with scripts/migrate_workspace.py)"
                    if os.path.exists(old) else "")
            raise ConfigError(f"{path} does not exist{hint}")
        return cls.from_yaml(path, section="excel")

    @staticmethod
    def _validate(cfg: dict) -> dict:
        cfg = copy.deepcopy(cfg)
        _req(cfg, "file", "")
        cases = _req(cfg, "cases", "")
        _req(cases, "sheets", "cases.")
        ccol = _req(cases, "columns", "cases.")
        _req(ccol, "id", "cases.columns.")
        _req(ccol, "status", "cases.columns.")
        _req(cases, "statuses", "cases.")
        cases.setdefault("header_row", 1)
        cases["fields"] = cases.get("fields") or {}
        unknown = [k for k in cases["fields"] if k not in CASE_MODEL_FIELDS]
        if unknown:
            raise ConfigError(f"configuration: cases.fields outside the case model {list(CASE_MODEL_FIELDS)}: {unknown}")
        cases["new_status"] = cases.get("new_status") or cases["statuses"][0]
        if cases["new_status"] not in cases["statuses"]:
            raise ConfigError(f"configuration: cases.new_status {cases['new_status']!r} outside cases.statuses")

        d = _req(cfg, "defects", "")
        _req(d, "sheet", "defects.")
        fmt = _req(d, "id_format", "defects.")
        if fmt.count("{n}") != 1:
            raise ConfigError("configuration: defects.id_format must contain '{n}' once")
        dcol = _req(d, "columns", "defects.")
        _req(dcol, "id", "defects.columns.")
        _req(dcol, "status", "defects.columns.")
        _req(d, "statuses", "defects.")
        _req(d, "initial_status", "defects.")
        d.setdefault("header_row", 1)
        for key in ("batch_fields", "defaults"):
            d[key] = d.get(key) or {}
        d["priorities"] = d.get("priorities") or []
        d["closed_statuses"] = d.get("closed_statuses") or []
        d["reported_by"] = d.get("reported_by") or REPORTED_BY_DEFAULT
        bad = [e for e in d["closed_statuses"] if e not in d["statuses"]]
        if bad:
            raise ConfigError(f"configuration: defects.closed_statuses outside defects.statuses: {bad}")
        if d["initial_status"] not in d["statuses"]:
            raise ConfigError("configuration: defects.initial_status outside defects.statuses")
        pre, _, post = fmt.partition("{n}")
        d["_id_re"] = re.compile("^" + re.escape(pre) + r"(\d+)" + re.escape(post) + "$")

        # optional blocks left blank (as in the template) count as absent
        v = cfg.get("active_view")
        if v and v.get("sheet"):
            v.setdefault("header_row", 1)
            v["columns"] = v.get("columns") or {}
            if not (v.get("summary") or {}).get("cell"):
                v["summary"] = None
        else:
            cfg["active_view"] = None

        r = cfg.get("summary")
        if r and r.get("sheet"):
            _req(r, "area_column", "summary.")
            r.setdefault("header_row", 1)
            r["status_columns"] = r.get("status_columns") or {}
            outside = [e for e in r["status_columns"] if e not in cases["statuses"]]
            if outside:
                raise ConfigError(f"configuration: summary.status_columns with statuses outside cases.statuses: "
                                  f"{outside}")
            if r.get("percent_column"):
                pend = _req(r, "not_run_status", "summary.")
                if pend not in cases["statuses"]:
                    raise ConfigError(f"configuration: summary.not_run_status {pend!r} outside cases.statuses")
                tol = r.get("percent_tolerance")
                r["percent_tolerance"] = 0.001 if tol in (None, "") else float(tol)
        else:
            cfg["summary"] = None

        defaults = {
            "notes": {"result": "Result: {result}", "new_defect": "New defect: {id}.",
                      "existing_defect": "Linked defect: {id}."},
            "batch_control": {"sheet": CONTROL_SHEET_DEFAULT},
            "recalc": {"mechanism": "libreoffice", "timeout": 60},
        }
        for block, vals in defaults.items():
            cfg[block] = cfg.get(block) or {}
            for k, val in vals.items():
                if cfg[block].get(k) in (None, ""):
                    cfg[block][k] = val
        return cfg

    @property
    def path(self) -> str:
        return os.path.join(self.base_dir, self.cfg["file"])

    # ------------------------------------------------------------------
    # Loading / saving
    # ------------------------------------------------------------------

    def load(self, path: str | None = None) -> openpyxl.Workbook:
        """Loads keeping formulas (data_only=False) -- what is wanted for
        EDITING. Calculated values are read in a separate copy."""
        self.wb = openpyxl.load_workbook(path or self.path, data_only=False)
        self._check_layout()
        return self.wb

    def _need_wb(self):
        if self.wb is None:
            raise RuntimeError("workbook not loaded -- call load() first")
        return self.wb

    @staticmethod
    def _headers(ws, row: int) -> dict:
        return {c.value: c.column for c in ws[row] if c.value not in (None, "")}

    def _check_layout(self) -> None:
        """Fails early if the configuration does not match the real workbook."""
        wb = self.wb
        problems = []
        cases, d = self.cfg["cases"], self.cfg["defects"]
        field_cols = set(cases["fields"].values())
        for sheet in set(cases["sheets"].values()):
            if sheet not in wb.sheetnames:
                problems.append(f"case sheet missing: '{sheet}'")
                continue
            h = self._headers(wb[sheet], cases["header_row"])
            field_cols -= set(h)
            # actual_result / notes are optional sheet by sheet
            for role in ("id", "status"):
                col = cases["columns"][role]
                if col not in h:
                    problems.append(f"'{sheet}': column '{col}' ({role}) missing")
        # a field's column may be missing from some sheets, never from all (a typo)
        if field_cols and all(sh in wb.sheetnames for sh in cases["sheets"].values()):
            problems += [f"cases.fields: column '{c}' is in no case sheet" for c in sorted(field_cols)]
        if d["sheet"] not in wb.sheetnames:
            problems.append(f"defects sheet missing: '{d['sheet']}'")
        else:
            h = self._headers(wb[d["sheet"]], d["header_row"])
            cols = [(k, c) for k, c in d["columns"].items()] + [(k, c) for k, c in d["batch_fields"].items()]
            for role, col in cols:
                if col and col not in h:
                    problems.append(f"'{d['sheet']}': column '{col}' ({role}) missing")
        v = self.cfg.get("active_view")
        if v and v["sheet"] not in wb.sheetnames:
            problems.append(f"active defects view sheet missing: '{v['sheet']}'")
        r = self.cfg.get("summary")
        if r:
            if r["sheet"] not in wb.sheetnames:
                problems.append(f"summary sheet missing: '{r['sheet']}'")
            else:
                h = self._headers(wb[r["sheet"]], r["header_row"])
                for role in ("area_column", "total_column", "percent_column"):
                    if r.get(role) and r[role] not in h:
                        problems.append(f"'{r['sheet']}': column '{r[role]}' ({role}) missing")
                if not any(r["status_columns"].get(e, e) in h for e in cases["statuses"]):
                    problems.append(f"'{r['sheet']}': no column named after a status (use status_columns)")
        if problems:
            raise ConfigError("configuration does not match the workbook:\n - " + "\n - ".join(problems))

    def save(self, path: str | None = None) -> str:
        out = path or self.path
        self._need_wb().save(out)
        return out

    def reviewed_path(self) -> str:
        root, ext = os.path.splitext(self.path)
        return f"{root}.REVIEWED{ext}"

    def save_if_clean(self) -> dict:
        """Saves over the original only if the cross-check returns zero
        problems; otherwise saves <name>.REVIEWED.xlsx and leaves the
        original untouched (technique-excel.md section 1, step 5)."""
        problems = self.cross_check()
        if problems:
            alt = self.save(self.reviewed_path())
            return {"saved_original": False, "file": alt, "problems": problems}
        return {"saved_original": True, "file": self.save(), "problems": []}

    def recalc(self, path: str | None = None) -> dict:
        r = self.cfg["recalc"]
        return recalc_safe(path or self.path, mechanism=r["mechanism"], timeout=r["timeout"])

    # ------------------------------------------------------------------
    # Reading
    # ------------------------------------------------------------------

    def _case_sheets(self):
        """(name, ws, headers) of each case sheet, in the configuration's order."""
        wb, cases = self._need_wb(), self.cfg["cases"]
        seen = []
        for sheet in cases["sheets"].values():
            if sheet not in seen:
                seen.append(sheet)
                ws = wb[sheet]
                yield sheet, ws, self._headers(ws, cases["header_row"])

    def _case_rows(self, ws, h):
        cases = self.cfg["cases"]
        idc = h[cases["columns"]["id"]]
        for r in range(cases["header_row"] + 1, ws.max_row + 1):
            v = ws.cell(r, idc).value
            if v not in (None, ""):
                yield r, v

    def sheet_for_case(self, case_id: str) -> str:
        for pref, sheet in self.cfg["cases"]["sheets"].items():
            if str(case_id).startswith(pref):
                return sheet
        raise KeyError(f"case ID without a configured prefix: {case_id}")

    def progress_summary(self) -> dict:
        """{sheet: {status: n}} -- direct count, not depending on formulas."""
        st = self.cfg["cases"]["columns"]["status"]
        out = {}
        for sheet, ws, h in self._case_sheets():
            cnt = Counter(ws.cell(r, h[st]).value for r, _ in self._case_rows(ws, h))
            out[sheet] = dict(cnt)
        return out

    def case_rows(self) -> dict:
        """{sheet: [row as dict by header, with '_row']} of every case, in the order of the sheets and of the
        rows (read only; what docs/qa-ui.md section 4 reads as the workspace's library)."""
        out = {}
        for sheet, ws, h in self._case_sheets():
            out[sheet] = [{**{name: ws.cell(r, c).value for name, c in h.items()}, "_row": r}
                          for r, _ in self._case_rows(ws, h)]
        return out

    def cases_with_status(self, status: str) -> dict:
        """{sheet: [row as dict, with '_row']} of the cases in a given status."""
        st = self.cfg["cases"]["columns"]["status"]
        out = {}
        for sheet, ws, h in self._case_sheets():
            rows = []
            for r, _ in self._case_rows(ws, h):
                if ws.cell(r, h[st]).value == status:
                    row = {name: ws.cell(r, c).value for name, c in h.items()}
                    row["_row"] = r
                    rows.append(row)
            out[sheet] = rows
        return out

    def _defect_rows(self):
        d = self.cfg["defects"]
        ws = self._need_wb()[d["sheet"]]
        h = self._headers(ws, d["header_row"])
        idc = h[d["columns"]["id"]]
        rows = [(r, ws.cell(r, idc).value) for r in range(d["header_row"] + 1, ws.max_row + 1)]
        return ws, h, [(r, v) for r, v in rows if v not in (None, "")]

    def _defect_number(self, defect_id) -> int | None:
        m = self.cfg["defects"]["_id_re"].match(str(defect_id).strip())
        return int(m.group(1)) if m else None

    def next_defect_number(self) -> int:
        _, _, rows = self._defect_rows()
        nums = [n for n in (self._defect_number(v) for _, v in rows) if n is not None]
        return max(nums, default=0) + 1

    # ------------------------------------------------------------------
    # Control of the batches applied
    # ------------------------------------------------------------------

    @staticmethod
    def _control_names(ws) -> list:
        """Header row of the control sheet, with the headers from before the translation read as the new ones."""
        return [_LEGACY_CONTROL_HEADERS.get(c.value, c.value) for c in ws[1]]

    def _control_ws(self, create: bool):
        wb, name = self._need_wb(), self.cfg["batch_control"]["sheet"]
        if name in wb.sheetnames:
            ws = wb[name]
            if create:  # sheets from earlier versions get the new columns (e.g. batch_id)
                existing = self._control_names(ws)
                for col in _CONTROL_HEADERS:
                    if col not in existing:
                        ws.cell(1, len(existing) + 1).value = col
                        existing.append(col)
            return ws
        if not create:
            return None
        ws = wb.create_sheet(name)
        ws.sheet_state = "hidden"
        ws.append(_CONTROL_HEADERS)
        return ws

    def applied_batches(self) -> list[dict]:
        ws = self._control_ws(create=False)
        if ws is None:
            return []
        head = self._control_names(ws)  # by the header row: old sheets have no batch_id
        return [dict(zip(head, row)) for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]

    def already_applied(self, batch: list, identity: str | None = None) -> dict | None:
        h = identity or batch_hash(batch)
        return next((x for x in self.applied_batches() if x["hash"] == h), None)

    # ------------------------------------------------------------------
    # Applying a batch
    # ------------------------------------------------------------------

    def validate_batch(self, batch: list) -> list[str]:
        """Read only. Problems that prevent applying the batch (empty = it can
        be applied): technique-excel.md section 4."""
        cases, d = self.cfg["cases"], self.cfg["defects"]
        if not isinstance(batch, list) or not batch:
            return ["empty batch"]
        existing = {str(v).strip() for _, ws, h in self._case_sheets() for _, v in self._case_rows(ws, h)}
        defects = {str(v).strip() for _, v in self._defect_rows()[2]}
        statuses, prios = cases["statuses"], d["priorities"]
        problems, seen = [], Counter()
        for i, item in enumerate(batch, 1):
            if not isinstance(item, dict):
                problems.append(f"item {i}: not an object")
                continue
            tc = str(item.get("tc") or "").strip()
            where = f"item {i} ({tc or 'no tc'})"
            if not tc:
                problems.append(f"{where}: 'tc' is missing")
            elif tc not in existing:
                problems.append(f"{where}: the case does not exist in the workbook")
            seen[tc] += 1
            if item.get("status") in (None, ""):
                problems.append(f"{where}: 'status' is missing")
            elif item["status"] not in statuses:
                problems.append(f"{where}: status {item['status']!r} outside {statuses}"
                                + (" -- resolve the Question before applying (methodology 8.5)"
                                   if str(item["status"]).lower().startswith("question") else ""))
            bug = item.get("bug")
            if bug not in (None, "") and str(bug).strip() not in defects:
                problems.append(f"{where}: bug {bug!r} does not exist in '{d['sheet']}'")
            new = item.get("new_defect")
            if new not in (None, {}):
                if not isinstance(new, dict):
                    problems.append(f"{where}: 'new_defect' must be an object or null")
                    continue
                if new.get("proposed", False) is not False:
                    problems.append(f"{where}: the new_defect is still a proposal (\"proposed\") -- the user approves "
                                    "it by removing \"proposed\" (or drops it) before the batch is applied "
                                    "(methodology 9.4)")
                if prios:
                    pr = new.get("priority", new.get("severity"))
                    if pr not in prios:
                        problems.append(f"{where}: priority of the new defect {pr!r} outside {prios}")
        dup = sorted(tc for tc, n in seen.items() if tc and n > 1)
        if dup:
            problems.append(f"case repeated in the batch: {dup}")
        return problems

    def apply_batch(self, batch: list, label: str | None = None, reported_by: str | None = None,
                    batch_id: str | None = None, identity: str | None = None) -> dict:
        """Applies the batch: status, actual result and notes per case, and
        creates the new defects. A batch already applied to this workbook (same
        content) is not applied again: returns {"applied": False, ...} without
        changing anything (technique-excel.md section 5). With `batch_id` (apply
        uses the file name), a batch with that id already applied but with a
        different content -- a batch edited after being applied -- is refused.
        `identity` replaces the content's hash as the batch's identity: a batch
        converted by scripts/migrate_workspace.py carries the hash of its
        original content (`translated_from`), which the workbook recorded.

        The whole batch is validated before anything is written
        (validate_batch); if it has problems, raises InvalidBatch and nothing
        changes."""
        self._need_wb()
        h_batch = identity or batch_hash(batch)
        previous = self.already_applied(batch, h_batch)
        if previous:
            return {"applied": False, "hash": h_batch, "new_defects": [],
                    "reason": f"batch already applied at {previous['applied_at']}"}
        if batch_id:
            same_id = next((x for x in self.applied_batches() if x.get("batch_id") == batch_id), None)
            if same_id:
                raise InvalidBatch([(
                    f"the batch '{batch_id}' was already applied at {same_id['applied_at']} with a DIFFERENT "
                    "content (it was edited after being applied). Applying it again would repeat notes and "
                    "defects. Put the correction in a new batch, with another name, with only the cases to "
                    "correct (technique-excel.md section 5)."
                )])

        label = label or f"[{self._now().date().isoformat()}]"
        cases, d, notes = self.cfg["cases"], self.cfg["defects"], self.cfg["notes"]
        reported_by = reported_by or d["reported_by"]
        ccol = cases["columns"]

        problems = self.validate_batch(batch)
        if problems:
            raise InvalidBatch(problems)

        # case -> (ws, headers, row)
        rowmap = {}
        for _sheet, ws, h in self._case_sheets():
            for r, v in self._case_rows(ws, h):
                rowmap.setdefault(str(v).strip(), (ws, h, r))

        ws_def, h_def, def_rows = self._defect_rows()
        last_def_row = max([r for r, _ in def_rows] + [d["header_row"]])
        next_num = self.next_defect_number()
        dcol = d["columns"]
        new_ids = []

        for item in batch:
            tc = str(item["tc"]).strip()
            ws, h, r = rowmap[tc]
            ws.cell(r, h[ccol["status"]]).value = item["status"]
            has_res = ccol.get("actual_result") in h
            if has_res:
                ws.cell(r, h[ccol["actual_result"]]).value = item.get("actual_result")

            new_defect = item.get("new_defect")
            def_id = None
            if new_defect:
                def_id = d["id_format"].replace("{n}", str(next_num))
                next_num += 1
                new_ids.append(def_id)

            if ccol.get("notes") in h:
                parts = []
                if not has_res:
                    parts.append(notes["result"].format(result=item.get("actual_result")))
                if def_id:
                    parts.append(notes["new_defect"].format(id=def_id))
                if item.get("bug") not in (None, ""):
                    parts.append(notes["existing_defect"].format(id=str(item["bug"]).strip()))
                if parts:
                    note = f"{label} " + " ".join(parts)
                    c = ws.cell(r, h[ccol["notes"]])
                    c.value = (str(c.value).rstrip() + "\n\n" + note).strip() if c.value else note

            if new_defect:
                last_def_row += 1
                vals = dict(d["defaults"])
                vals.update(new_defect)

                def put(col, value, row=last_def_row):
                    if col:
                        ws_def.cell(row, h_def[col]).value = value

                put(dcol["id"], def_id)
                put(dcol.get("case"), vals.get("related_tc", tc))
                for key, col in d["batch_fields"].items():
                    put(col, vals.get(key))
                put(dcol.get("severity"), vals.get("severity"))
                put(dcol.get("priority"), vals.get("priority", vals.get("severity")))
                put(dcol["status"], d["initial_status"])
                put(dcol.get("reported_by"), reported_by)
                put(dcol.get("report_date"), self._now().date().isoformat())

        ctl = self._control_ws(create=True)
        record = {"hash": h_batch, "applied_at": self._now().isoformat(timespec="seconds"),
                  "label": label, "items": len(batch), "new_defects": ", ".join(new_ids),
                  "batch_id": batch_id}
        ctl.append([record.get(name) for name in self._control_names(ctl)])
        return {"applied": True, "hash": h_batch, "new_defects": new_ids}

    # ------------------------------------------------------------------
    # Adding approved cases (technique-excel.md section 4.1)
    # ------------------------------------------------------------------

    @staticmethod
    def _title_key(value) -> str:
        return " ".join(str(value or "").split()).casefold()

    @staticmethod
    def case_cell_text(field: str, value):
        """The text a model field gets in its cell: steps one per line, lists joined."""
        if value in (None, "", []):
            return None
        if field == "steps" and isinstance(value, list):
            lines = []
            for n, step in enumerate(value, 1):
                line = f"{n}. {str(step.get('action') or '').strip()}"
                if step.get("data") not in (None, ""):
                    line += f" [{step['data']}]"
                if step.get("expected") not in (None, ""):
                    line += f" -> {step['expected']}"
                lines.append(line)
            return "\n".join(lines)
        if isinstance(value, list):
            return ", ".join(str(v) for v in value)
        return value

    def validate_cases(self, cases: list) -> list[str]:
        """Read only. Problems that prevent adding the cases (empty = they can
        be added): technique-excel.md section 4.1."""
        ccfg = self.cfg["cases"]
        if "title" not in ccfg["fields"]:
            return [("cases.fields has no 'title' column: add-cases needs it to find duplicates "
                     "(technique-excel.md section 4.1)")]
        if not isinstance(cases, list) or not cases:
            return ["no cases"]
        existing = {str(v).strip() for _, ws, h in self._case_sheets() for _, v in self._case_rows(ws, h)}
        titles: dict[str, dict[str, str]] = {}  # sheet -> {title key: case id}
        for sheet, ws, h in self._case_sheets():
            tcol = h.get(ccfg["fields"]["title"])
            titles[sheet] = {} if tcol is None else {
                self._title_key(ws.cell(r, tcol).value): str(v).strip()
                for r, v in self._case_rows(ws, h) if ws.cell(r, tcol).value not in (None, "")}
        problems, seen = [], Counter()
        for i, case in enumerate(cases, 1):
            if not isinstance(case, dict):
                problems.append(f"case {i}: not an object")
                continue
            cid = str(case.get("id") or "").strip()
            where = f"case {i} ({cid or 'no id'})"
            sheet = None
            if not cid:
                problems.append(f"{where}: 'id' is missing")
            else:
                seen[cid] += 1
                try:
                    sheet = self.sheet_for_case(cid)
                except KeyError:
                    problems.append(f"{where}: the id has no prefix of cases.sheets {list(ccfg['sheets'])}")
                if cid in existing:
                    problems.append(f"{where}: the case already exists in the workbook")
            title = self._title_key(case.get("title"))
            if not title:
                problems.append(f"{where}: 'title' is missing")
            elif sheet:
                other = titles[sheet].get(title)
                if other and other != cid:
                    problems.append(f"{where}: candidate duplicate of {other} (same title in '{sheet}'): remove the "
                                    "case from the file if it is the same case (methodology 10)")
                else:
                    titles[sheet][title] = cid
            steps = case.get("steps")
            if steps is not None and not (isinstance(steps, list) and all(isinstance(x, dict) for x in steps)):
                problems.append(f"{where}: 'steps' must be a list of {{action, data, expected}}")
        dup = sorted(c for c, n in seen.items() if n > 1)
        if dup:
            problems.append(f"case id repeated in the file: {dup}")
        return problems

    def plan_cases(self, cases: list) -> list[dict]:
        """Read only, for valid cases: per case {"id", "sheet", "row", "cells":
        {column: value}, "not_written": [fields]}, in the file's order."""
        ccfg = self.cfg["cases"]
        sheets = {sheet: (ws, h) for sheet, ws, h in self._case_sheets()}
        last = {sheet: max([r for r, _ in self._case_rows(ws, h)] + [ccfg["header_row"]])
                for sheet, (ws, h) in sheets.items()}
        out = []
        for case in cases:
            cid = str(case["id"]).strip()
            sheet = self.sheet_for_case(cid)
            _, h = sheets[sheet]
            last[sheet] += 1
            cells = {ccfg["columns"]["id"]: cid, ccfg["columns"]["status"]: ccfg["new_status"]}
            not_written = []
            for field in CASE_MODEL_FIELDS:
                value = self.case_cell_text(field, case.get(field))
                if value is None:
                    continue
                col = ccfg["fields"].get(field)
                if col and col in h:
                    cells[col] = value
                else:
                    not_written.append(field)
            out.append({"id": cid, "sheet": sheet, "row": last[sheet], "cells": cells, "not_written": not_written})
        return out

    @staticmethod
    def _copy_row_format(ws, h: dict, src: int, dst: int, written: set) -> None:
        """Style of row `src` on row `dst`; formulas of the columns not written translated; data
        validations that covered `src` extended to `dst`. The row height is never set (section 3)."""
        for col in h.values():
            a, b = ws.cell(src, col), ws.cell(dst, col)
            if a.has_style:
                b.font, b.fill, b.border = copy.copy(a.font), copy.copy(a.fill), copy.copy(a.border)
                b.alignment, b.protection = copy.copy(a.alignment), copy.copy(a.protection)
                b.number_format = a.number_format
            if col not in written and isinstance(a.value, str) and a.value.startswith("="):
                b.value = Translator(a.value, origin=a.coordinate).translate_formula(b.coordinate)
        for dv in ws.data_validations.dataValidation:
            for col in h.values():
                a, b = ws.cell(src, col).coordinate, ws.cell(dst, col).coordinate
                if a in dv.sqref and b not in dv.sqref:
                    dv.sqref.add(b)

    def add_cases(self, cases: list, label: str | None = None, batch_id: str | None = None) -> dict:
        """Adds approved cases (technique-excel.md section 4.1): each in the row
        after the last case of its sheet, with the previous row's style, in the
        status cases.new_status. Recorded in the control sheet like a batch: the
        same cases again change nothing ({"added": False}); the same batch_id
        with another content is refused. Everything is validated before
        writing (InvalidBatch)."""
        self._need_wb()
        h_cases = batch_hash([{"add_cases": cases}])
        previous = next((x for x in self.applied_batches() if x["hash"] == h_cases), None)
        if previous:
            return {"added": False, "hash": h_cases, "cases": [],
                    "reason": f"cases already added at {previous['applied_at']}"}
        if batch_id:
            same_id = next((x for x in self.applied_batches() if x.get("batch_id") == batch_id), None)
            if same_id:
                raise InvalidBatch([(
                    f"the cases file '{batch_id}' was already added at {same_id['applied_at']} with a DIFFERENT "
                    "content (it was edited after being added). New cases go in a new file, with another name "
                    "(technique-excel.md section 4.1)."
                )])
        problems = self.validate_cases(cases)
        if problems:
            raise InvalidBatch(problems)

        plan = self.plan_cases(cases)
        sheets = {sheet: (ws, h) for sheet, ws, h in self._case_sheets()}
        for p in plan:
            ws, h = sheets[p["sheet"]]
            written = {h[c] for c in p["cells"]}
            if p["row"] - 1 > self.cfg["cases"]["header_row"]:
                self._copy_row_format(ws, h, p["row"] - 1, p["row"], written)
            for col, value in p["cells"].items():
                ws.cell(p["row"], h[col]).value = value

        label = label or f"[{self._now().date().isoformat()}]"
        ctl = self._control_ws(create=True)
        record = {"hash": h_cases, "applied_at": self._now().isoformat(timespec="seconds"),
                  "label": label, "items": len(cases), "new_defects": "", "batch_id": batch_id}
        ctl.append([record.get(name) for name in self._control_names(ctl)])
        return {"added": True, "hash": h_cases, "cases": plan}

    def update_defect_note(self, defect_id: str, note: str, column: str,
                           new_status: str | None = None) -> None:
        """Appends a note to a text column of an existing defect, without
        erasing what was there. Changes the status only if new_status is
        passed explicitly."""
        d = self.cfg["defects"]
        ws, h, rows = self._defect_rows()
        if column not in h:
            raise KeyError(f"column '{column}' does not exist in '{d['sheet']}'")
        for r, v in rows:
            if str(v).strip() == defect_id:
                c = ws.cell(r, h[column])
                c.value = (str(c.value or "").rstrip() + "\n\n" + note).strip()
                if new_status:
                    if new_status not in d["statuses"]:
                        raise ValueError(f"invalid defect status: {new_status}")
                    ws.cell(r, h[d["columns"]["status"]]).value = new_status
                return
        raise KeyError(f"defect not found: {defect_id}")

    # ------------------------------------------------------------------
    # Active defects view
    # ------------------------------------------------------------------

    def open_defects(self, sort_by: str | None = None) -> dict:
        """Read only. {"open": [row as dict], "closed": n}: the defects whose
        status is not in closed_statuses, sorted by priority, by `sort_by`
        (defects column, optional) and by the defect's number."""
        d = self.cfg["defects"]
        ws_def, h_def, def_rows = self._defect_rows()
        st_col = h_def[d["columns"]["status"]]
        closed_set = set(d["closed_statuses"])

        rows, closed = [], 0
        for r, _ in def_rows:
            if ws_def.cell(r, st_col).value in closed_set:
                closed += 1
                continue
            rows.append({name: ws_def.cell(r, c).value for name, c in h_def.items()})

        prio_col, prios = d["columns"].get("priority"), d["priorities"]

        def key(row):
            p = row.get(prio_col) if prio_col else None
            pk = prios.index(p) if p in prios else len(prios)
            n = self._defect_number(row.get(d["columns"]["id"]))
            sec = str(row.get(sort_by) or "") if sort_by else ""
            return (pk, sec, n if n is not None else float("inf"))

        rows.sort(key=key)
        return {"open": rows, "closed": closed}

    def regenerate_active_defects(self) -> dict:
        """Rebuilds the data rows of the active defects view from the defects
        sheet (without the closed statuses), sorted by priority, by the
        configured secondary criterion and by the defect's number. Keeps what
        is above the first data row and reapplies the style of the original
        first data row."""
        v = self.cfg.get("active_view")
        if not v:
            return {"total_open": None, "reason": "active_view not configured"}
        d = self.cfg["defects"]
        od = self.open_defects(v.get("sort_by"))
        rows, closed = od["open"], od["closed"]

        ws = self._need_wb()[v["sheet"]]
        hv = self._headers(ws, v["header_row"])
        first = v["header_row"] + 1
        styles = None
        if ws.cell(first, 1).value not in (None, ""):
            styles = {}
            for col in hv.values():
                c = ws.cell(first, col)
                styles[col] = (copy.copy(c.font), copy.copy(c.fill), copy.copy(c.border),
                               copy.copy(c.alignment), c.number_format)
        if ws.max_row >= first:
            ws.delete_rows(first, ws.max_row - first + 1)

        for i, row in enumerate(rows):
            target = first + i
            for head, col in hv.items():
                c = ws.cell(target, col)
                c.value = row.get(v["columns"].get(head, head))
                if styles:
                    c.font, c.fill, c.border, c.alignment, c.number_format = styles[col]

        sev_col = d["columns"].get("severity")
        sev = Counter(row.get(sev_col) for row in rows) if sev_col else Counter()
        summary = v.get("summary")
        if summary:
            ws[summary["cell"]].value = summary["template"].format_map(
                {"open": len(rows), "closed": closed, "severity": sev, "date": self._now().date()}
            )
        return {"total_open": len(rows), "by_severity": dict(sev), "closed": closed}

    # ------------------------------------------------------------------
    # Cross-check (methodology section 11)
    # ------------------------------------------------------------------

    def cross_check(self) -> list[str]:
        """Returns a list of problems as text (empty = 0 problems): case and
        defect statuses outside the valid list, duplicated case and defect
        IDs, and defects linked to a case that does not exist."""
        cases, d = self.cfg["cases"], self.cfg["defects"]
        case_statuses, def_statuses = set(cases["statuses"]), set(d["statuses"])
        problems = []
        case_ids = Counter()
        for sheet, ws, h in self._case_sheets():
            sc = h[cases["columns"]["status"]]
            for r, v in self._case_rows(ws, h):
                case_ids[str(v).strip()] += 1
                st = ws.cell(r, sc).value
                if st not in case_statuses:
                    problems.append(f"{sheet} row {r} ({v}): invalid case status {st!r}")
        dup = sorted(k for k, n in case_ids.items() if n > 1)
        if dup:
            problems.append(f"duplicated case ID: {dup}")

        ws_def, h_def, def_rows = self._defect_rows()
        sc = h_def[d["columns"]["status"]]
        case_col = d["columns"].get("case")
        seen = Counter()
        for r, v in def_rows:
            seen[str(v).strip()] += 1
            st = ws_def.cell(r, sc).value
            if st not in def_statuses:
                problems.append(f"{d['sheet']} row {r} ({v}): invalid defect status {st!r}")
            if case_col:
                ref = ws_def.cell(r, h_def[case_col]).value
                for tc in self._case_refs(ref):
                    if tc not in case_ids:
                        problems.append(f"{d['sheet']} row {r} ({v}): related case does not exist {tc!r}")
        dup = sorted(k for k, n in seen.items() if n > 1)
        if dup:
            problems.append(f"duplicated defect ID: {dup}")
        return problems

    def check_summary(self, path: str | None = None) -> dict:
        """Read only. Compares the CALCULATED values of the summary sheet
        (excel.summary) with the direct count of the statuses, both read from
        the file saved at `path` (by default the workbook): per area, per
        status, in the area's total and in the total row (technique-excel.md
        section 5). Returns {"checked", "problems", "reason"}; checked is
        False without a summary or without calculated values (recalculate
        first)."""
        r = self.cfg.get("summary")
        if not r:
            return {"checked": False, "problems": [], "reason": "summary not configured"}
        wb = openpyxl.load_workbook(path or self.path, data_only=True)
        reader = copy.copy(self)
        reader.wb = wb
        direct = {area: Counter(c) for area, c in reader.progress_summary().items()}
        total = sum(direct.values(), Counter())

        ws = wb[r["sheet"]]
        h = self._headers(ws, r["header_row"])
        col_area = h[r["area_column"]]
        col_st = {e: h[r["status_columns"].get(e, e)] for e in self.cfg["cases"]["statuses"]
                  if r["status_columns"].get(e, e) in h}
        col_tot = h.get(r["total_column"]) if r.get("total_column") else None
        # % run (methodology 16.3): (total - not run) / total; N/A and Blocked count
        col_pct = h.get(r["percent_column"]) if r.get("percent_column") else None
        pend, tol_pct = r.get("not_run_status"), r.get("percent_tolerance")

        problems, seen, cells, empty = [], set(), 0, 0
        total_row_seen = False
        for row in range(r["header_row"] + 1, ws.max_row + 1):
            area = ws.cell(row, col_area).value
            area = "" if area is None else str(area).strip()
            if r.get("total_row") and area == r["total_row"]:
                target, name, total_row_seen = total, "total row", True
            elif area in direct:
                target, name = direct[area], f"area '{area}'"
                seen.add(area)
            else:
                continue  # titles, notes and empty rows are not areas
            compare = [(e, c, target.get(e, 0)) for e, c in col_st.items()]
            if col_tot:
                compare.append(("total", col_tot, sum(target.values())))
            if col_pct:
                n = sum(target.values())
                compare.append(("% run", col_pct, (n - target.get(pend, 0)) / n if n else 0.0))
            for caption, col, expected in compare:
                cell = ws.cell(row, col)
                cells += 1
                if cell.value is None:
                    empty += 1
                    continue
                try:
                    v = float(cell.value)
                    if col == col_pct:  # fraction (0.2857) or percentage points (28.57)
                        equal = min(abs(v - expected * 100), abs(v - expected) * 100) <= tol_pct
                    else:
                        equal = v == expected
                except (TypeError, ValueError):
                    equal = False
                if not equal:
                    rule = (f"(total - {pend}) / total gives {expected:.4%}" if col == col_pct
                            else f"the direct count gives {expected}")
                    problems.append(f"{r['sheet']}!{cell.coordinate} ({name}, {caption}): the summary says "
                                    f"{cell.value!r}, {rule}")
        if cells and empty == cells:
            return {"checked": False, "problems": [],
                    "reason": "the summary has no calculated values -- recalculate first"}
        if empty:
            problems.append(f"{r['sheet']}: {empty} cell(s) without a calculated value")
        problems += [f"{r['sheet']}: the area '{a}' is not in the summary" for a in direct if a not in seen]
        if r.get("total_row") and not total_row_seen:
            problems.append(f"{r['sheet']}: total row '{r['total_row']}' not found")
        return {"checked": True, "problems": problems, "reason": None}

    def _case_refs(self, value) -> list[str]:
        """Case IDs named in a cell (there may be several, separated by comma,
        semicolon, space or line break). Only what starts with a configured
        case prefix counts as a reference."""
        if value in (None, ""):
            return []
        prefixes = tuple(self.cfg["cases"]["sheets"].keys())
        return [t for t in re.split(r"[\s,;/]+", str(value)) if t.startswith(prefixes)]
