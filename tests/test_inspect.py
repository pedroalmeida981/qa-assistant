"""Inspecting a workbook for the onboarding (adapters/excel/inspect.py;
technique-excel.md section 0): detects cases, defects and the active view only by
the relations between the sheets, with no predefined names.

python -m unittest discover -s tests   (from the root of qa-assistant)
"""

import os
import sys
import tempfile
import unittest

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from adapters.excel.inspect import propose

STATUSES = ["To run", "Passed", "Failed"]


def workbook(path, defects, view, title_in_view=True):
    """Cases TC-SHOP-001..006; `defects` and `view` are lists of (id, status, case)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Shop"
    ws.append(["ID", "Title", "State", "Actual Result"])
    for i in range(1, 7):
        ws.append([f"TC-SHOP-{i:03d}", f"Case {i}", STATUSES[i % 3], None])
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"')
    dv.add("C2:C200")
    ws.add_data_validation(dv)

    d = wb.create_sheet("Register")
    d.append(["ID", "Summary", "Situation", "Case", "Gravity"])
    for bid, st, case in defects:
        d.append([bid, f"Problem {bid}", st, case, "High" if bid.endswith(("1", "3")) else "Low"])

    v = wb.create_sheet("Open")
    if title_in_view:
        v.append(["Open defects (generated)"])
        v.append([None])
    v.append(["ID", "Summary", "Situation", "Case", "Gravity"])
    for bid, st, case in view:
        v.append([bid, f"Problem {bid}", st, case, "High" if bid.endswith(("1", "3")) else "Low"])
    wb.save(path)


class Inspect(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.tmp.name, "plan.xlsx")

    def tearDown(self):
        self.tmp.cleanup()

    def test_cases_detected_by_the_validation_list(self):
        defs = [("D-1", "Open", "TC-SHOP-001"), ("D-2", "Closed", "TC-SHOP-002"),
                ("D-3", "Open", "TC-SHOP-003")]
        workbook(self.path, defs, [d for d in defs if d[1] == "Open"])
        cfg, _ = propose(self.path)
        self.assertEqual(cfg["cases"]["sheets"], {"TC-SHOP-": "Shop"})
        self.assertEqual(cfg["cases"]["columns"]["id"], "ID")
        self.assertEqual(cfg["cases"]["columns"]["status"], "State")
        self.assertEqual(cfg["cases"]["statuses"], STATUSES)

    def test_view_with_fewer_ids_than_the_register(self):
        defs = [("D-1", "Open", "TC-SHOP-001"), ("D-2", "Closed", "TC-SHOP-002"),
                ("D-3", "Open", "TC-SHOP-003")]
        workbook(self.path, defs, [d for d in defs if d[1] == "Open"])
        cfg, _ = propose(self.path)
        self.assertEqual(cfg["defects"]["sheet"], "Register")
        self.assertEqual(cfg["defects"]["id_format"], "D-{n}")
        self.assertEqual(cfg["defects"]["columns"]["case"], "Case")
        self.assertEqual(cfg["active_view"]["sheet"], "Open")
        self.assertEqual(cfg["active_view"]["header_row"], 3)

    def test_every_defect_open_view_equal_to_the_register(self):
        # new project: no closed defect, the view has the same IDs as the register.
        # Before, each one was a "view" of the other and the defects sheet got no proposal.
        defs = [("D-1", "Open", "TC-SHOP-001"), ("D-2", "Open", "TC-SHOP-002"),
                ("D-3", "Open", "TC-SHOP-004")]
        workbook(self.path, defs, defs)
        cfg, notes = propose(self.path)
        self.assertEqual(cfg["defects"]["sheet"], "Register")
        self.assertEqual(cfg["defects"]["columns"]["case"], "Case")
        self.assertEqual(cfg["active_view"]["sheet"], "Open")
        self.assertFalse([n for n in notes if "candidate views" in n])
        # the status column, with a single value, is among the options; the related case column is not
        st = cfg["defects"]["columns"]["status"]
        self.assertIn("'Situation'", st)
        self.assertNotIn("'Case'", st)

    def test_same_ids_without_a_tie_break_is_left_to_confirm(self):
        defs = [("D-1", "Open", "TC-SHOP-001"), ("D-2", "Open", "TC-SHOP-002")]
        workbook(self.path, defs, defs, title_in_view=False)
        cfg, _ = propose(self.path)
        self.assertTrue(str(cfg["defects"]["sheet"]).startswith("CONFIRM"))
        self.assertIn("Register", cfg["defects"]["sheet"])
        self.assertIn("Open", cfg["defects"]["sheet"])

    def test_the_adapters_control_sheet_is_ignored_also_with_the_old_name(self):
        defs = [("D-1", "Open", "TC-SHOP-001"), ("D-2", "Closed", "TC-SHOP-002")]
        workbook(self.path, defs, [defs[0]])
        wb = openpyxl.load_workbook(self.path)
        wb.create_sheet("_qa_lotes").append(["hash", "aplicado_em", "etiqueta", "itens", "defeitos_novos", "lote_id"])
        wb.save(self.path)
        cfg, notes = propose(self.path)
        self.assertTrue(any("_qa_lotes" in n and "ignored" in n for n in notes), notes)
        self.assertEqual(cfg["defects"]["sheet"], "Register")

    def test_never_saves_the_workbook(self):
        defs = [("D-1", "Open", "TC-SHOP-001")]
        workbook(self.path, defs, defs)
        before = os.path.getmtime(self.path), os.path.getsize(self.path)
        propose(self.path)
        self.assertEqual((os.path.getmtime(self.path), os.path.getsize(self.path)), before)


if __name__ == "__main__":
    unittest.main()
