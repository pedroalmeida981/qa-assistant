# -*- coding: utf-8 -*-
"""
Gera o workbook FICTICIO do teste golden (plano, tarefa 1.4).

O layout corresponde a adapters/excel/exemplo.yaml: se o exemplo mudar e
deixar de bater com esta fixture, o golden falha -- e deve falhar.
O resultado e sempre o mesmo (sem datas nem valores aleatorios).

    python tests/golden/gerar_fixture.py <destino.xlsx>
"""
from __future__ import annotations

import sys

import openpyxl
from openpyxl.styles import Font, PatternFill

CASE_SHEETS = {
    # sheet: (prefixo, tem coluna "Actual Result", casos [(sufixo, titulo, estado, notas)])
    "Loja": ("CT-LOJA-", True, [
        ("001", "Add item to cart", "Pass", None),
        ("002", "Checkout with card", "To Do", None),
        ("003", "Apply discount code", "To Do", None),
        ("004", "Order total with tax", "Fail", "see BUG-1"),
    ]),
    "Conta": ("CT-CONTA-", True, [
        ("001", "Save profile", "To Do", "pre-existing note"),
        ("002", "Login with valid user", "Pass", None),
        ("003", "Reset password", "Blocked", "mail server down"),
        ("004", "Delete account", "To Do", None),
    ]),
    # sem coluna "Actual Result": o resultado vai para as notas
    "Admin": ("CT-ADMIN-", False, [
        ("001", "Export report", "To Do", None),
        ("002", "Legacy import", "N/A", "out of scope"),
        ("003", "Bulk edit users", "To Do", None),
        ("004", "Audit log view", "Pass", None),
    ]),
}

DEFECT_HEADERS = ["Key", "Summary", "Description", "Component", "Severity",
                  "Priority", "Status", "Test Case", "Reporter", "Created"]
DEFECTS = [
    ["BUG-1", "Checkout total wrong", "Tax not added to total", "Loja", "High", "P2", "New",
     "CT-LOJA-004", "tester-a", "2026-01-10"],
    ["BUG-2", "Login timeout", "Session expires after 1 min", "Conta", "Medium", "P3", "Resolved",
     "CT-CONTA-002", "tester-b", "2026-01-11"],
    ["BUG-3", "Admin export slow", "Export takes > 5 min", "Admin", "Low", "P3", "In Progress",
     "CT-ADMIN-001", "tester-a", "2026-01-12"],
]

VIEW_HEADERS = ["Bug", "Prio", "Sev", "Status", "Component", "Title", "Case"]
STATES = ["To Do", "Pass", "Fail", "Blocked", "N/A"]


def gerar(path: str) -> str:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sheet, (pre, has_result, cases) in CASE_SHEETS.items():
        ws = wb.create_sheet(sheet)
        headers = ["ID", "Title", "Status"] + (["Actual Result"] if has_result else []) + ["Notes"]
        ws.append(headers)
        for suf, title, state, notes in cases:
            ws.append([f"{pre}{suf}", title, state] + ([None] if has_result else []) + [notes])

    ws = wb.create_sheet("Defects")
    ws.append(DEFECT_HEADERS)
    for row in DEFECTS:
        ws.append(row)

    # vista de ativos desatualizada de proposito (so BUG-1), com estilo na 1.a linha de dados
    ws = wb.create_sheet("Open Defects")
    ws["A1"] = "Open defects (fictitious)"
    ws["A2"] = "summary not generated yet"
    ws.append(VIEW_HEADERS)
    ws.append(["BUG-1", "P2", "High", "New", "Loja", "Checkout total wrong", "CT-LOJA-004"])
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFEEEE")

    ws = wb.create_sheet("Summary")
    ws.append(["Sheet", "Total"] + STATES)
    for sheet in CASE_SHEETS:
        ws.append([sheet, f"=COUNTA('{sheet}'!A2:A100)"]
                  + [f"=COUNTIF('{sheet}'!C2:C100,\"{s}\")" for s in STATES])
    n = len(CASE_SHEETS)
    ws.append(["Total"] + [f"=SUM({col}2:{col}{n + 1})" for col in "BCDEFG"])

    wb.save(path)
    return path


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("uso: python tests/golden/gerar_fixture.py <destino.xlsx>")
    print(gerar(sys.argv[1]))
