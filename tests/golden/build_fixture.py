"""
Builds the FICTIONAL workbook of the golden test (plan, task 1.4).

The layout matches adapters/excel/example.yaml: if the example changes and no
longer matches this fixture, the golden fails -- and it should.
The result is always the same (no dates nor random values).

    python tests/golden/build_fixture.py <target.xlsx>
"""
from __future__ import annotations

import sys

import openpyxl
from openpyxl.styles import Font, PatternFill

CASE_SHEETS = {
    # sheet: (prefix, has an "Actual Result" column, cases [(suffix, title, status, notes)])
    "Shop": ("TC-SHOP-", True, [
        ("001", "Add item to cart", "Pass", None),
        ("002", "Checkout with card", "To Do", None),
        ("003", "Apply discount code", "To Do", None),
        ("004", "Order total with tax", "Fail", "see BUG-1"),
    ]),
    "Account": ("TC-ACCOUNT-", True, [
        ("001", "Save profile", "To Do", "pre-existing note"),
        ("002", "Login with valid user", "Pass", None),
        ("003", "Reset password", "Blocked", "mail server down"),
        ("004", "Delete account", "To Do", None),
    ]),
    # no "Actual Result" column: the result goes to the notes
    "Admin": ("TC-ADMIN-", False, [
        ("001", "Export report", "To Do", None),
        ("002", "Legacy import", "N/A", "out of scope"),
        ("003", "Bulk edit users", "To Do", None),
        ("004", "Audit log view", "Pass", None),
    ]),
}

DEFECT_HEADERS = ["Key", "Summary", "Description", "Component", "Severity",
                  "Priority", "Status", "Test Case", "Reporter", "Created"]
DEFECTS = [
    ["BUG-1", "Checkout total wrong", "Tax not added to total", "Shop", "High", "P2", "New",
     "TC-SHOP-004", "tester-a", "2026-01-10"],
    ["BUG-2", "Login timeout", "Session expires after 1 min", "Account", "Medium", "P3", "Resolved",
     "TC-ACCOUNT-002", "tester-b", "2026-01-11"],
    ["BUG-3", "Admin export slow", "Export takes > 5 min", "Admin", "Low", "P3", "In Progress",
     "TC-ADMIN-001", "tester-a", "2026-01-12"],
]

VIEW_HEADERS = ["Bug", "Prio", "Sev", "Status", "Component", "Title", "Case"]
STATES = ["To Do", "Pass", "Fail", "Blocked", "N/A"]


def build(path: str) -> str:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sheet, (pre, has_result, cases) in CASE_SHEETS.items():
        ws = wb.create_sheet(sheet)
        headers = ["ID", "Title", "Status"] + (["Actual Result"] if has_result else []) + ["Notes"]
        ws.append(headers)
        for suf, title, state, notes in cases:
            ws.append([f"{pre}{suf}", title, state] + ([None] if has_result else []) + [notes])

    ws = wb.create_sheet("Defects")
    ws.append(DEFECT_HEADERS)
    for row in DEFECTS:
        ws.append(row)

    # active view out of date on purpose (only BUG-1), with a style on the 1st data row
    ws = wb.create_sheet("Open Defects")
    ws["A1"] = "Open defects (fictitious)"
    ws["A2"] = "summary not generated yet"
    ws.append(VIEW_HEADERS)
    ws.append(["BUG-1", "P2", "High", "New", "Shop", "Checkout total wrong", "TC-SHOP-004"])
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFEEEE")

    ws = wb.create_sheet("Summary")
    ws.append(["Sheet", "Total"] + STATES)
    for sheet in CASE_SHEETS:
        ws.append([sheet, f"=COUNTA('{sheet}'!A2:A100)"]
                  + [f"=COUNTIF('{sheet}'!C2:C100,\"{s}\")" for s in STATES])
    n = len(CASE_SHEETS)
    ws.append(["Total"] + [f"=SUM({col}2:{col}{n + 1})" for col in "BCDEFG"])

    wb.save(path)
    return path


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: python tests/golden/build_fixture.py <target.xlsx>")
    print(build(sys.argv[1]))
