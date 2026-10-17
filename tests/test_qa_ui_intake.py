"""qa-ui intake (docs/qa-ui.md section 4; plan 12.1): pasted text, files and the workspace's library, the notes kept
apart and the flags. Fictional data only (Demo Shop); the library is the Excel adapter's golden fixture.

python -m unittest discover -s tests   (from the root of qa-assistant)
"""

import base64
import io
import json
import os
import sys
import unittest
import zipfile

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path[:0] = [ROOT, os.path.join(ROOT, "scripts"), os.path.dirname(os.path.abspath(__file__))]

import qa_ui_kit as kit
from common import read_project

from adapters.excel import ExcelAdapter
from qa_ui import intake as it

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def b64(data: bytes) -> str:
    return base64.b64encode(data).decode("ascii")


def docx(body: str, prolog: str = "") -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", f'<?xml version="1.0"?>{prolog}<w:document xmlns:w="{W}"><w:body>{body}'
                                        "</w:body></w:document>")
    return buf.getvalue()


def p(text: str) -> str:
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def row(*cells: str) -> str:
    return "<w:tr>" + "".join(f"<w:tc>{p(c)}</w:tc>" for c in cells) + "</w:tr>"


class Text(unittest.TestCase):
    def test_one_case_per_line_with_each_arrow(self):
        cases, notes = it.parse_text("CHK-011 Pay with a valid card → order created\n"
                                     "CHK-012 Pay with a refused card -> an error message is shown\n"
                                     "CHK-013: Apply an expired coupon => the coupon is refused\n")
        self.assertEqual([(c["id"], c["title"], c["expected"]) for c in cases],
                         [("CHK-011", "Pay with a valid card", "order created"),
                          ("CHK-012", "Pay with a refused card", "an error message is shown"),
                          ("CHK-013", "Apply an expired coupon", "the coupon is refused")])
        self.assertEqual(notes, [])
        self.assertTrue(all(c["flags"] == [] and c["source"] == "text" for c in cases))

    def test_a_line_without_an_id_is_a_note_and_never_a_case(self):
        cases, notes = it.parse_text("The checkout has a new payment form.\n\n- CHK-011 Pay → order created\n"
                                     "## Release notes\n4712 Something\nCheckout must be fast\n")
        self.assertEqual([c["id"] for c in cases], ["CHK-011"])
        self.assertEqual(notes, ["The checkout has a new payment form.", "Release notes", "4712 Something",
                                 "Checkout must be fast"])

    def test_an_id_of_the_library_is_an_id_whatever_its_form(self):
        cases, notes = it.parse_text("4712 Pay with a valid card → order created", known=frozenset({"4712"}))
        self.assertEqual([(c["id"], c["title"]) for c in cases], [("4712", "Pay with a valid card")])
        self.assertEqual(notes, [])

    def test_cells_by_tab_or_pipe_in_the_default_order(self):
        cases, _ = it.parse_text("CHK-011\tPay with a valid card\torder created\n"
                                 "CHK-012 | Pay with a refused card | an error message is shown\n")
        self.assertEqual([(c["id"], c["expected"]) for c in cases],
                         [("CHK-011", "order created"), ("CHK-012", "an error message is shown")])

    def test_a_markdown_table_with_its_header_names_the_columns(self):
        cases, notes = it.parse_text("| Expected Result | Case ID | Title | Role | Platform |\n"
                                     "|---|:--|--:|---|---|\n"
                                     "| order created | CHK-011 | Pay with a valid card | customer | android |\n"
                                     "| | not an id | a row that is not a case | | |\n")
        self.assertEqual(len(cases), 1)
        c = cases[0]
        self.assertEqual((c["id"], c["title"], c["expected"], c["role"], c["platform"]),
                         ("CHK-011", "Pay with a valid card", "order created", "customer", "android"))
        self.assertEqual(notes, ["not an id | a row that is not a case"])

    def test_flags(self):
        cases, _ = it.parse_text("LOG-001 Login with a wrong password → generic error\n"
                                 "LOG-002 Login with a valid password\n"
                                 "LOG-003 Login → works correctly\n"
                                 "LOG-004 Login → the error message is shown correctly under the password field\n")
        self.assertEqual([c["flags"] for c in cases],
                         [["vague-expected"], ["no-expected"], ["vague-expected"], []])
        several = it.new_case({"id": "LOG-005", "title": "Login",
                               "expected": "1. The page opens\n2. The name is shown"}, "text")
        self.assertEqual(several["flags"], ["several-expected"])

    def test_the_same_id_twice_keeps_the_first_and_flags_it(self):
        a, _ = it.parse_text("CHK-011 First → one\nCHK-012 Other → two")
        b, _ = it.parse_text("CHK-011 Second → three")
        cases, dropped = it.merge([a, b])
        self.assertEqual([(c["id"], c["title"]) for c in cases], [("CHK-011", "First"), ("CHK-012", "Other")])
        self.assertEqual(cases[0]["flags"], ["duplicate-id"])
        self.assertEqual(dropped, ["CHK-011"])


class Files(unittest.TestCase):
    def test_csv_with_headers_in_any_order_and_a_semicolon(self):
        data = "Title;ID;Expected Result;Steps\nPay with a valid card;CHK-011;order created;\"1. Add\n2. Pay\"\n"
        cases, _ = it.read_file("cases.csv", b64(data.encode("utf-8")))
        self.assertEqual((cases[0]["id"], cases[0]["title"], cases[0]["steps"]),
                         ("CHK-011", "Pay with a valid card", "1. Add\n2. Pay"))

    def test_csv_without_headers_and_tsv(self):
        cases, _ = it.read_file("cases.csv", b64(b"CHK-011,Pay with a valid card,order created\n"))
        self.assertEqual((cases[0]["id"], cases[0]["expected"]), ("CHK-011", "order created"))
        cases, _ = it.read_file("cases.tsv", b64(b"id\ttitle\nCHK-012\tPay, with a comma\n"))
        self.assertEqual((cases[0]["id"], cases[0]["title"]), ("CHK-012", "Pay, with a comma"))

    def test_txt_and_md_are_read_as_pasted_text(self):
        for name in ("cases.txt", "cases.md"):
            cases, notes = it.read_file(name, b64("Notes first\nCHK-011 Pay → order created\n".encode()))
            self.assertEqual(([c["id"] for c in cases], notes, cases[0]["source"]),
                             (["CHK-011"], ["Notes first"], "file"))

    def test_xlsx_every_sheet_with_a_header_and_the_sheet_as_the_area(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Checkout"
        ws.append(["Demo Shop — test plan"])
        ws.append(["ID", "Title", "Expected Result", "Status"])
        ws.append(["CHK-011", "Pay with a valid card", "order created", "Pass"])
        ws.append([None, "a row without an id", None, None])
        other = wb.create_sheet("Login")
        other.append(["Case", "Name"])
        other.append(["LOG-001", "Login with a wrong password"])
        wb.create_sheet("Readme").append(["nothing", "here"])
        buf = io.BytesIO()
        wb.save(buf)
        cases, notes = it.read_file("plan.xlsx", b64(buf.getvalue()))
        self.assertEqual([(c["id"], c["area"], c["flags"]) for c in cases],
                         [("CHK-011", "Checkout", []), ("LOG-001", "Login", ["no-expected"])])
        self.assertEqual(notes, [])

    def test_docx_paragraphs_and_tables_in_the_documents_order(self):
        data = docx(p("The checkout has a new payment form.") + p("CHK-011 Pay with a valid card → order created")
                    + "<w:tbl>" + row("ID", "Title", "Expected") + row("CHK-012", "Pay with a refused card",
                                                                       "an error message is shown") + "</w:tbl>"
                    + p("LOG-001 Login with a wrong password"))
        cases, notes = it.read_file("spec.docx", b64(data))
        self.assertEqual([c["id"] for c in cases], ["CHK-011", "CHK-012", "LOG-001"])
        self.assertEqual(notes, ["The checkout has a new payment form."])

    def test_what_is_refused_says_why(self):
        for name, data, text in (
                ("spec.pdf", b"%PDF-1.7", ".pdf is not read"),
                ("cases.json", b"{}", ".json is not read"),
                ("plan.xlsx", b"not a workbook", "not a readable .xlsx"),
                ("spec.docx", b"not a zip", "not a readable .docx"),
                ("spec.docx", docx("", "<!DOCTYPE x [<!ENTITY a 'b'>]>"), "entities"),
                ("big.txt", b"x" * (it.MAX_FILE + 1), "more than 5 MB")):
            with self.subTest(name=name, text=text):
                with self.assertRaises(it.IntakeError) as cm:
                    it.read_file(name, b64(data))
                self.assertIn(text, str(cm.exception))
        with self.assertRaises(it.IntakeError):
            it.read_file("cases.txt", "not base64!")


class Library(unittest.TestCase):
    def setUp(self):
        self.tmp = kit.workspace()
        self.ws = self.tmp.name
        self.proj = read_project(self.ws)

    def tearDown(self):
        self.tmp.cleanup()

    def test_the_workbook_through_the_adapter_read_only(self):
        before = os.path.getmtime(os.path.join(self.ws, kit.EXCEL["file"]))
        out = it.intake(self.ws, self.proj, library=True)
        self.assertEqual(out["library"]["source"], "workbook")
        self.assertIsNone(out["library"]["error"])
        self.assertEqual(out["library"]["count"], 12)
        first = out["cases"][0]
        self.assertEqual((first["id"], first["title"], first["area"], first["last_status"], first["in_library"]),
                         ("TC-SHOP-001", "Add item to cart", "Shop", "Pass", True))
        self.assertEqual(os.path.getmtime(os.path.join(self.ws, kit.EXCEL["file"])), before)

    def test_adapter_case_rows_is_every_case_in_order(self):
        ad = ExcelAdapter.from_project(self.ws)
        ad.load()
        rows = ad.case_rows()
        self.assertEqual(list(rows), ["Shop", "Account", "Admin"])
        self.assertEqual([r["ID"] for r in rows["Shop"]], [f"TC-SHOP-{n:03d}" for n in range(1, 5)])
        self.assertEqual(rows["Shop"][0]["_row"], 2)

    def test_pasted_cases_are_marked_and_filled_from_the_library(self):
        out = it.intake(self.ws, self.proj, text="TC-SHOP-001 Add item\nNEW-1 A new idea → something observable")
        known, new = out["cases"]
        self.assertEqual((known["in_library"], known["title"], known["area"], known["last_status"]),
                         (True, "Add item", "Shop", "Pass"))
        self.assertEqual((new["in_library"], new["flags"]), (False, ["not-in-library"]))

    def test_text_and_library_together_do_not_duplicate(self):
        out = it.intake(self.ws, self.proj, text="TC-SHOP-001 Add item → the cart shows 1", library=True)
        self.assertEqual(len(out["cases"]), 12)
        self.assertEqual(out["dropped"], [])
        self.assertEqual(out["cases"][0]["flags"], [])

    def test_a_workbook_that_cannot_be_read_is_said_and_nothing_is_marked(self):
        os.remove(os.path.join(self.ws, kit.EXCEL["file"]))
        out = it.intake(self.ws, self.proj, text="NEW-1 A new idea → something observable")
        self.assertIn("could not be read", out["library"]["error"])
        self.assertEqual((out["cases"][0]["in_library"], out["cases"][0]["flags"]), (None, []))

    def _headers(self, sheet: str, renamed: dict, added: dict | None = None) -> None:
        """Renames headers of a sheet of the fixture workbook and adds a column (header -> one value per row)."""
        path = os.path.join(self.ws, kit.EXCEL["file"])
        wb = openpyxl.load_workbook(path)
        ws = wb[sheet]
        for cell in ws[1]:
            if cell.value in renamed:
                cell.value = renamed[cell.value]
        for header, value in (added or {}).items():
            col = ws.max_column + 1
            ws.cell(1, col, header)
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, 1).value:
                    ws.cell(r, col, value)
        wb.save(path)

    def test_headers_in_another_language_are_read_without_a_mapping(self):
        proj = dict(self.proj, excel=dict(self.proj["excel"], cases=dict(self.proj["excel"]["cases"])))
        proj["excel"]["cases"].pop("fields", None)
        kit.write_project(self.ws, proj)
        self._headers("Shop", {"Title": "Título"}, {"Résultat attendu": "the cart shows the item",
                                                    "PRÉ-CONDIÇÕES": "a customer session"})
        lib = it.load_library(self.ws, read_project(self.ws))
        first = lib["cases"][0]
        self.assertEqual((first["title"], first["expected"], first["preconditions"]),
                         ("Add item to cart", "the cart shows the item", "a customer session"))

    def test_the_projects_mapping_wins_over_a_recognised_header(self):
        proj = read_project(self.ws)
        proj["excel"]["cases"]["fields"] = {"title": "Notes"}
        kit.write_project(self.ws, proj)
        ad = ExcelAdapter.from_project(self.ws)
        ad.load()
        row = ad.case_rows()["Shop"][0]
        lib = it.load_library(self.ws, read_project(self.ws))
        self.assertEqual(lib["cases"][0]["title"], it.text_of(row["Notes"]))

    def test_cases_without_title_or_expected_are_warned_with_the_cause(self):
        lib = it.load_library(self.ws, self.proj)
        self.assertIn("12 of 12 library case(s) cannot be judged", lib["warning"])
        self.assertIn("12 without an expected result (TC-SHOP-001, TC-SHOP-002, TC-SHOP-003, …)", lib["warning"])
        self.assertIn("excel.cases.fields", lib["warning"])
        self.assertEqual(it.intake(self.ws, self.proj, library=True)["library"]["warning"], lib["warning"])
        self._headers("Shop", {}, {"Expected Result": "x"})
        self._headers("Account", {}, {"Expected Result": "x"})
        self._headers("Admin", {}, {"Expected Result": "x"})
        self.assertIsNone(it.load_library(self.ws, self.proj)["warning"])

    def test_headers_compare_without_case_accents_or_separators(self):
        self.assertEqual(it.norm("  Pré-Condições_ "), "pre condicoes")
        for header, field in (("Título", "title"), ("ERWARTETES ERGEBNIS", "expected"), ("Résultat attendu",
                              "expected"), ("Pasos", "steps"), ("Perfil", "role"), ("Prioridade", "priority")):
            self.assertEqual(it._HEADER_OF.get(it.norm(header)), field, header)

    def test_tracker_cases_json_as_import_cases_wrote_it(self):
        proj = dict(self.proj, tracker="jira")
        out = it.intake(self.ws, proj, library=True)
        self.assertIn("import-cases", out["library"]["error"])
        os.makedirs(os.path.join(self.ws, "tracker"))
        doc = {"tracker": "jira", "imported_at": "2026-10-01T10:00:00", "flags": [], "cases": [
            {"id": "DS-140", "key": "DS-140", "title": "Pay with a valid card", "area": "Checkout",
             "steps": [{"action": "Add an item", "data": None, "expected": None},
                       {"action": "Pay", "data": "a valid card", "expected": "order created"}],
             "expected": "order created", "role": "customer", "platform": None, "last_status": "Pass"}]}
        with open(os.path.join(self.ws, "tracker", "cases.json"), "w", encoding="utf-8") as f:
            json.dump(doc, f)
        out = it.intake(self.ws, proj, library=True)
        case = out["cases"][0]
        self.assertEqual(out["library"]["detail"], "1 case(s), imported at 2026-10-01T10:00:00")
        self.assertEqual((case["id"], case["steps"], case["last_status"], case["flags"]),
                         ("DS-140", "1. Add an item\n2. Pay [a valid card] → order created", "Pass", []))


if __name__ == "__main__":
    unittest.main()
