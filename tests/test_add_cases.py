"""Adding approved cases to an Excel workbook (technique-excel.md section 4.1; plan 10.6).

python -m unittest discover -s tests   (from the root of qa-assistant)
Does not need LibreOffice: never recalculates.
"""

import contextlib
import copy
import io
import json
import os
import sys
import tempfile
import unittest
from typing import ClassVar

import openpyxl
import yaml
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path[:0] = [ROOT, os.path.join(ROOT, "tests", "golden")]

from build_fixture import build

from adapters.excel import ExcelAdapter, InvalidBatch
from adapters.excel import add_cases as cli
from adapters.excel.base import ConfigError

with open(os.path.join(ROOT, "adapters", "excel", "example.yaml"), encoding="utf-8") as _f:
    CFG = yaml.safe_load(_f)


def case(cid="TC-SHOP-005", title="Remove item from cart", **extra):
    return {"id": cid, "title": title, **extra}


class Base(unittest.TestCase):
    cfg_changes: ClassVar[dict] = {}

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.xlsx = os.path.join(self.tmp.name, CFG["file"])
        build(self.xlsx)
        self.prepare(self.xlsx)
        cfg = copy.deepcopy(CFG)
        cfg["cases"].update(copy.deepcopy(self.cfg_changes))
        self.ad = ExcelAdapter.from_dict(cfg, base_dir=self.tmp.name)
        self.ad.load()

    def prepare(self, path):
        """Hook: changes the fixture before it is loaded."""

    def tearDown(self):
        self.tmp.cleanup()

    def ids(self, sheet):
        ws = self.ad.wb[sheet]
        return [ws.cell(r, 1).value for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]

    def refused(self, cases, fragment):
        before = {s: self.ids(s) for s in ("Shop", "Account", "Admin")}
        with self.assertRaises(InvalidBatch) as cm:
            self.ad.add_cases(cases)
        self.assertIn(fragment, str(cm.exception))
        self.assertEqual({s: self.ids(s) for s in before}, before)
        self.assertEqual(self.ad.applied_batches(), [])


class AddCases(Base):
    def test_cases_go_after_the_last_case_of_their_sheet(self):
        r = self.ad.add_cases([case(), case("TC-ADMIN-005", "Archive users"), case("TC-SHOP-006", "Empty cart")],
                              label="[x]", batch_id="cases_01")
        self.assertTrue(r["added"])
        self.assertEqual(self.ids("Shop")[-2:], ["TC-SHOP-005", "TC-SHOP-006"])
        self.assertEqual(self.ids("Admin")[-1], "TC-ADMIN-005")
        ws = self.ad.wb["Shop"]
        self.assertEqual([ws.cell(6, c).value for c in (1, 2, 3)], ["TC-SHOP-005", "Remove item from cart", "To Do"])
        self.assertEqual([p["row"] for p in r["cases"]], [6, 6, 7])
        self.assertEqual(self.ad.progress_summary()["Shop"]["To Do"], 4)
        self.assertEqual(self.ad.cross_check(), [])
        rec = self.ad.applied_batches()[0]
        self.assertEqual((rec["batch_id"], rec["items"], rec["label"]), ("cases_01", 3, "[x]"))

    def test_the_same_file_again_changes_nothing(self):
        self.ad.add_cases([case()], batch_id="cases_01")
        r = self.ad.add_cases([case()], batch_id="cases_01")
        self.assertFalse(r["added"])
        self.assertEqual(self.ids("Shop").count("TC-SHOP-005"), 1)
        self.assertEqual(len(self.ad.applied_batches()), 1)

    def test_a_file_edited_after_being_added_is_refused(self):
        self.ad.add_cases([case()], batch_id="cases_01")
        with self.assertRaises(InvalidBatch) as cm:
            self.ad.add_cases([case("TC-SHOP-006", "Empty cart")], batch_id="cases_01")
        self.assertIn("DIFFERENT", str(cm.exception))

    def test_a_results_batch_can_then_run_the_new_case(self):
        self.ad.add_cases([case()], batch_id="cases_01")
        self.ad.apply_batch([{"tc": "TC-SHOP-005", "status": "Pass", "actual_result": "removed"}], batch_id="b1")
        self.assertEqual(self.ad.wb["Shop"].cell(6, 3).value, "Pass")

    def test_fields_without_a_column_are_listed_not_refused(self):
        r = self.ad.add_cases([case(feature="Cart", technique="State transition")])
        self.assertEqual(r["cases"][0]["not_written"], ["feature", "technique"])

    def test_the_status_of_a_new_case(self):
        self.assertEqual(self.ad.plan_cases([case()])[0]["cells"]["Status"], "To Do")


class Refused(Base):
    def test_an_id_that_exists(self):
        self.refused([case("TC-SHOP-002", "Another title")], "already exists in the workbook")

    def test_an_id_without_a_configured_prefix(self):
        self.refused([case("TC-NONE-001")], "no prefix of cases.sheets")

    def test_an_id_repeated_in_the_file(self):
        self.refused([case(), case(title="Other")], "case id repeated in the file")

    def test_a_title_already_in_the_sheet_is_a_candidate_duplicate(self):
        self.refused([case(title="  add ITEM to   cart ")], "candidate duplicate of TC-SHOP-001")

    def test_the_same_title_twice_in_the_file(self):
        self.refused([case(), case("TC-SHOP-006")], "candidate duplicate of TC-SHOP-005")

    def test_the_same_title_in_another_sheet_is_not_a_duplicate(self):
        self.assertEqual(self.ad.validate_cases([case("TC-ADMIN-005", "Add item to cart")]), [])

    def test_without_title(self):
        self.refused([case(title="")], "'title' is missing")

    def test_without_id(self):
        self.refused([{"title": "x"}], "'id' is missing")

    def test_steps_must_be_objects(self):
        self.refused([case(steps=["open the cart"])], "'steps' must be a list")

    def test_empty_file(self):
        self.refused([], "no cases")

    def test_without_a_title_column(self):
        del self.ad.cfg["cases"]["fields"]["title"]
        self.refused([case()], "no 'title' column")

    def test_one_bad_case_refuses_the_whole_file(self):
        self.refused([case(), case("TC-SHOP-002", "x")], "already exists")


class CellText(Base):
    cfg_changes: ClassVar[dict] = {"fields": {"title": "Title", "steps": "Notes", "role": "Actual Result"}}

    def test_steps_one_per_line_and_lists_joined(self):
        steps = [{"action": "Open the cart", "data": "", "expected": ""},
                 {"action": "Remove the item", "data": "SKU-1", "expected": "The cart is empty"}]
        p = self.ad.plan_cases([case(steps=steps, role=["Guest", "Customer"])])[0]
        self.assertEqual(p["cells"]["Notes"], "1. Open the cart\n2. Remove the item [SKU-1] -> The cart is empty")
        self.assertEqual(p["cells"]["Actual Result"], "Guest, Customer")

    def test_a_column_missing_from_one_sheet(self):
        # Admin has no "Actual Result": role is not written there, and listed
        p = self.ad.plan_cases([case("TC-ADMIN-005", "Archive users", role="Admin")])[0]
        self.assertEqual(p["not_written"], ["role"])


class RowFormat(Base):
    def prepare(self, path):
        wb = openpyxl.load_workbook(path)
        ws = wb["Shop"]
        ws["E1"] = "Key"
        for r in range(2, 6):
            ws.cell(r, 5).value = f'=A{r}&"-"&C{r}'
        ws["B5"].font = Font(bold=True)
        ws["B5"].fill = PatternFill("solid", fgColor="FFEEEE")
        ws["C5"].number_format = "@"
        dv = DataValidation(type="list", formula1='"To Do,Pass,Fail,Blocked,N/A"')
        dv.add("C2:C5")
        ws.add_data_validation(dv)
        wb.save(path)

    def test_style_formula_and_validation_follow_the_previous_row(self):
        self.ad.add_cases([case()])
        ws = self.ad.wb["Shop"]
        self.assertTrue(ws["B6"].font.b)
        self.assertEqual(ws["B6"].fill.fgColor.rgb, ws["B5"].fill.fgColor.rgb)
        self.assertEqual(ws["C6"].number_format, "@")
        self.assertEqual(ws["E6"].value, '=A6&"-"&C6')
        self.assertIn("C6", ws.data_validations.dataValidation[0].sqref)
        self.assertIsNone(ws.row_dimensions[6].height)

    def test_the_formula_is_kept_through_save_and_load(self):
        self.ad.add_cases([case()])
        self.assertTrue(self.ad.save_if_clean()["saved_original"])
        ws = openpyxl.load_workbook(self.xlsx)["Shop"]
        self.assertEqual((ws["A6"].value, ws["E6"].value), ("TC-SHOP-005", '=A6&"-"&C6'))


class Configuration(unittest.TestCase):
    def cfg(self, **cases):
        c = copy.deepcopy(CFG)
        c["cases"].update(cases)
        return c

    def test_a_field_outside_the_model(self):
        with self.assertRaises(ConfigError) as cm:
            ExcelAdapter.from_dict(self.cfg(fields={"title": "Title", "summary": "Title"}))
        self.assertIn("outside the case model", str(cm.exception))

    def test_new_status_outside_the_statuses(self):
        with self.assertRaises(ConfigError):
            ExcelAdapter.from_dict(self.cfg(new_status="Open"))

    def test_a_column_in_no_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            build(os.path.join(tmp, CFG["file"]))
            ad = ExcelAdapter.from_dict(self.cfg(fields={"title": "Titel"}), base_dir=tmp)
            with self.assertRaises(ConfigError) as cm:
                ad.load()
            self.assertIn("column 'Titel' is in no case sheet", str(cm.exception))

    def test_without_fields_the_adapter_still_loads(self):
        c = copy.deepcopy(CFG)
        del c["cases"]["fields"]
        self.assertEqual(ExcelAdapter.from_dict(c).cfg["cases"]["fields"], {})


class CommandLine(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.ws = self.tmp.name
        build(os.path.join(self.ws, CFG["file"]))
        with open(os.path.join(self.ws, "project.yaml"), "w", encoding="utf-8") as f:
            yaml.safe_dump({"tracker": "excel", "excel": CFG}, f)
        os.makedirs(os.path.join(self.ws, "cases"))

    def tearDown(self):
        self.tmp.cleanup()

    def write(self, name, cases):
        with open(os.path.join(self.ws, "cases", name), "w", encoding="utf-8") as f:
            json.dump({"cases": cases}, f)
        return f"cases/{name}"

    def run_cli(self, *args):
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            code = cli.main([self.ws, *args, "--no-recalc", "--now", "2026-10-04T10:00:00"])
        return code, out.getvalue() + err.getvalue()

    def shop_ids(self):
        ws = openpyxl.load_workbook(os.path.join(self.ws, CFG["file"]))["Shop"]
        return [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]

    def test_add_then_again(self):
        f = self.write("cases_20261004_01.json", [case()])
        code, out = self.run_cli(f)
        self.assertEqual(code, 0, out)
        self.assertIn("TC-SHOP-005 -> 'Shop' row 6", out)
        self.assertEqual(self.shop_ids()[-1], "TC-SHOP-005")
        code, out = self.run_cli(f)
        self.assertEqual(code, 0)
        self.assertIn("Nothing added", out)
        ad = ExcelAdapter.from_project(self.ws)
        ad.load()
        self.assertEqual(ad.applied_batches()[0]["label"], "[2026-10-04, cases_20261004_01]")

    def test_dry_run_writes_nothing(self):
        before = self.shop_ids()
        code, out = self.run_cli(self.write("c.json", [case()]), "--dry-run")
        self.assertEqual(code, 0, out)
        self.assertIn("Dry run", out)
        self.assertEqual(self.shop_ids(), before)

    def test_invalid_file_is_exit_4(self):
        code, out = self.run_cli(self.write("c.json", [case("TC-SHOP-001", "x")]))
        self.assertEqual(code, 4)
        self.assertIn("NOTHING ADDED", out)
        code, _ = self.run_cli(self.write("d.json", [case("TC-SHOP-001", "x")]), "--dry-run")
        self.assertEqual(code, 4)

    def test_wrong_format(self):
        path = os.path.join(self.ws, "cases", "bad.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump([case()], f)
        self.assertEqual(self.run_cli("cases/bad.json")[0], 4)


if __name__ == "__main__":
    unittest.main()
