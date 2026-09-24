"""Inspecao de um workbook para o onboarding (adapters/excel/inspecionar.py;
tecnica-excel.md secao 0): deteta casos, defeitos e vista de ativos so pelas
relacoes entre as sheets, sem nomes pre-definidos.

python -m unittest discover -s tests   (a partir da raiz do qa-assistant)
"""

import os
import sys
import tempfile
import unittest

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from adapters.excel.inspecionar import propose

ESTADOS = ["Por executar", "Passou", "Falhou"]


def workbook(path, defeitos, vista, titulo_na_vista=True):
    """Casos CT-LOJA-001..006; `defeitos` e `vista` sao listas de (id, estado, caso)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Loja"
    ws.append(["ID", "Titulo", "Estado", "Resultado Obtido"])
    for i in range(1, 7):
        ws.append([f"CT-LOJA-{i:03d}", f"Caso {i}", ESTADOS[i % 3], None])
    dv = DataValidation(type="list", formula1='"' + ",".join(ESTADOS) + '"')
    dv.add("C2:C200")
    ws.add_data_validation(dv)

    d = wb.create_sheet("Registo")
    d.append(["ID", "Resumo", "Situacao", "Caso", "Gravidade"])
    for bid, est, caso in defeitos:
        d.append([bid, f"Problema {bid}", est, caso, "Alta" if bid.endswith(("1", "3")) else "Baixa"])

    v = wb.create_sheet("Abertos")
    if titulo_na_vista:
        v.append(["Defeitos em aberto (gerado)"])
        v.append([None])
    v.append(["ID", "Resumo", "Situacao", "Caso", "Gravidade"])
    for bid, est, caso in vista:
        v.append([bid, f"Problema {bid}", est, caso, "Alta" if bid.endswith(("1", "3")) else "Baixa"])
    wb.save(path)


class Inspecionar(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.tmp.name, "plano.xlsx")

    def tearDown(self):
        self.tmp.cleanup()

    def test_casos_detetados_pela_lista_de_validacao(self):
        defs = [("D-1", "Aberto", "CT-LOJA-001"), ("D-2", "Fechado", "CT-LOJA-002"),
                ("D-3", "Aberto", "CT-LOJA-003")]
        workbook(self.path, defs, [d for d in defs if d[1] == "Aberto"])
        cfg, _ = propose(self.path)
        self.assertEqual(cfg["casos"]["sheets"], {"CT-LOJA-": "Loja"})
        self.assertEqual(cfg["casos"]["colunas"]["id"], "ID")
        self.assertEqual(cfg["casos"]["colunas"]["estado"], "Estado")
        self.assertEqual(cfg["casos"]["estados"], ESTADOS)

    def test_vista_com_menos_ids_que_o_registo(self):
        defs = [("D-1", "Aberto", "CT-LOJA-001"), ("D-2", "Fechado", "CT-LOJA-002"),
                ("D-3", "Aberto", "CT-LOJA-003")]
        workbook(self.path, defs, [d for d in defs if d[1] == "Aberto"])
        cfg, _ = propose(self.path)
        self.assertEqual(cfg["defeitos"]["sheet"], "Registo")
        self.assertEqual(cfg["defeitos"]["formato_id"], "D-{n}")
        self.assertEqual(cfg["defeitos"]["colunas"]["caso"], "Caso")
        self.assertEqual(cfg["vista_ativos"]["sheet"], "Abertos")
        self.assertEqual(cfg["vista_ativos"]["linha_cabecalho"], 3)

    def test_todos_os_defeitos_abertos_vista_igual_ao_registo(self):
        # projeto novo: nenhum defeito fechado, a vista tem os mesmos IDs que o registo.
        # Antes cada uma era "vista" da outra e a sheet de defeitos ficava sem proposta.
        defs = [("D-1", "Aberto", "CT-LOJA-001"), ("D-2", "Aberto", "CT-LOJA-002"),
                ("D-3", "Aberto", "CT-LOJA-004")]
        workbook(self.path, defs, defs)
        cfg, notas = propose(self.path)
        self.assertEqual(cfg["defeitos"]["sheet"], "Registo")
        self.assertEqual(cfg["defeitos"]["colunas"]["caso"], "Caso")
        self.assertEqual(cfg["vista_ativos"]["sheet"], "Abertos")
        self.assertFalse([n for n in notas if "vistas candidatas" in n])
        # a coluna de estado, com um so valor, aparece nas opcoes; a do caso relacionado nao
        est = cfg["defeitos"]["colunas"]["estado"]
        self.assertIn("'Situacao'", est)
        self.assertNotIn("'Caso'", est)

    def test_mesmos_ids_sem_desempate_fica_para_confirmar(self):
        defs = [("D-1", "Aberto", "CT-LOJA-001"), ("D-2", "Aberto", "CT-LOJA-002")]
        workbook(self.path, defs, defs, titulo_na_vista=False)
        cfg, _ = propose(self.path)
        self.assertTrue(str(cfg["defeitos"]["sheet"]).startswith("CONFIRMAR"))
        self.assertIn("Registo", cfg["defeitos"]["sheet"])
        self.assertIn("Abertos", cfg["defeitos"]["sheet"])

    def test_nunca_grava_o_workbook(self):
        defs = [("D-1", "Aberto", "CT-LOJA-001")]
        workbook(self.path, defs, defs)
        antes = os.path.getmtime(self.path), os.path.getsize(self.path)
        propose(self.path)
        self.assertEqual((os.path.getmtime(self.path), os.path.getsize(self.path)), antes)


if __name__ == "__main__":
    unittest.main()
